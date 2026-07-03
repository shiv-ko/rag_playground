"""Extract style/mark information from DOCX and PPTX files."""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from docx import Document as DocxDocument
from pptx import Presentation


ROOT = Path(__file__).resolve().parent.parent
SHARE_ROOT = ROOT / "data" / "raw" / "share" / "共有ドライブ"
PROJECT_ROOT = SHARE_ROOT / "プロジェクト"
INTERNAL_ROOT = SHARE_ROOT / "社内管理"
ARTIFACTS = ROOT / "artifacts"


def write_jsonl(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as f:
        for row in rows:
            f.write(json.dumps(row, ensure_ascii=False) + "\n")


def project_and_section(path: Path) -> tuple[str | None, str]:
    try:
        rel = path.relative_to(PROJECT_ROOT)
    except ValueError:
        return None, "社内管理"
    parts = rel.parts
    return parts[0], parts[1] if len(parts) > 1 else ""


def color_value(color: Any) -> str | None:
    if color is None:
        return None
    rgb = getattr(color, "rgb", None)
    if rgb:
        return str(rgb)
    theme_color = getattr(color, "theme_color", None)
    if theme_color:
        return f"THEME:{theme_color}"
    return None


def docx_highlight(run: Any) -> str | None:
    value = getattr(run.font, "highlight_color", None)
    return str(value) if value else None


def base(path: Path) -> dict[str, Any]:
    project, section = project_and_section(path)
    return {
        "source_path": str(path.relative_to(ROOT)),
        "project_name": project,
        "section": section,
        "file_name": path.name,
        "extension": path.suffix.lower(),
    }


def extract_docx(path: Path) -> tuple[list[dict[str, Any]], list[dict[str, str]]]:
    rows: list[dict[str, Any]] = []
    failures: list[dict[str, str]] = []
    meta = base(path)
    try:
        doc = DocxDocument(path)
    except Exception as exc:  # noqa: BLE001
        return rows, [{"source_path": meta["source_path"], "error": repr(exc)}]

    for para_idx, para in enumerate(doc.paragraphs):
        for run_idx, run in enumerate(para.runs):
            text = run.text
            if not text.strip():
                continue
            font_color = color_value(run.font.color)
            row = {
                **meta,
                "unit_type": "paragraph_run",
                "paragraph_index": para_idx,
                "run_index": run_idx,
                "text": text,
                "bold": bool(run.bold),
                "italic": bool(run.italic),
                "underline": bool(run.underline),
                "font_color": font_color,
                "highlight_color": docx_highlight(run),
            }
            if any([row["bold"], row["italic"], row["underline"], row["font_color"], row["highlight_color"]]):
                rows.append(row)

    for table_idx, table in enumerate(doc.tables):
        for row_idx, table_row in enumerate(table.rows):
            for col_idx, cell in enumerate(table_row.cells):
                for para_idx, para in enumerate(cell.paragraphs):
                    for run_idx, run in enumerate(para.runs):
                        text = run.text
                        if not text.strip():
                            continue
                        font_color = color_value(run.font.color)
                        row = {
                            **meta,
                            "unit_type": "table_run",
                            "table_index": table_idx,
                            "row_index": row_idx,
                            "column_index": col_idx,
                            "paragraph_index": para_idx,
                            "run_index": run_idx,
                            "text": text,
                            "bold": bool(run.bold),
                            "italic": bool(run.italic),
                            "underline": bool(run.underline),
                            "font_color": font_color,
                            "highlight_color": docx_highlight(run),
                        }
                        if any([row["bold"], row["italic"], row["underline"], row["font_color"], row["highlight_color"]]):
                            rows.append(row)

    return rows, failures


def extract_pptx(path: Path) -> tuple[list[dict[str, Any]], list[dict[str, str]]]:
    rows: list[dict[str, Any]] = []
    meta = base(path)
    try:
        prs = Presentation(path)
    except Exception as exc:  # noqa: BLE001
        return rows, [{"source_path": meta["source_path"], "error": repr(exc)}]

    for slide_idx, slide in enumerate(prs.slides, start=1):
        for shape_idx, shape in enumerate(slide.shapes):
            if not getattr(shape, "has_text_frame", False):
                continue
            for para_idx, para in enumerate(shape.text_frame.paragraphs):
                for run_idx, run in enumerate(para.runs):
                    text = run.text
                    if not text.strip():
                        continue
                    font = run.font
                    font_color = color_value(font.color)
                    fill_color = None
                    try:
                        fill_color = color_value(shape.fill.fore_color)
                    except Exception:  # noqa: BLE001
                        fill_color = None
                    row = {
                        **meta,
                        "unit_type": "slide_run",
                        "slide_number": slide_idx,
                        "shape_index": shape_idx,
                        "paragraph_index": para_idx,
                        "run_index": run_idx,
                        "text": text,
                        "bold": bool(font.bold),
                        "italic": bool(font.italic),
                        "underline": bool(font.underline),
                        "font_color": font_color,
                        "fill_color": fill_color,
                    }
                    if any([row["bold"], row["italic"], row["underline"], row["font_color"], row["fill_color"]]):
                        rows.append(row)
    return rows, []


def target_files() -> list[Path]:
    roots = [PROJECT_ROOT, INTERNAL_ROOT]
    files: list[Path] = []
    for root in roots:
        files.extend(p for p in root.rglob("*.docx") if not p.name.startswith("~$"))
        files.extend(p for p in root.rglob("*.pptx") if not p.name.startswith("~$"))
    return sorted(files)


def main() -> None:
    marks: list[dict[str, Any]] = []
    failures: list[dict[str, str]] = []
    for path in target_files():
        print(f"extracting {path.relative_to(ROOT)}", flush=True)
        if path.suffix.lower() == ".docx":
            rows, errs = extract_docx(path)
        else:
            rows, errs = extract_pptx(path)
        marks.extend(rows)
        failures.extend(errs)
    write_jsonl(ARTIFACTS / "office_marks.jsonl", marks)
    write_jsonl(ARTIFACTS / "office_mark_failures.jsonl", failures)
    print(f"files={len(target_files())}")
    print(f"marks={len(marks)}")
    print(f"failures={len(failures)}")


if __name__ == "__main__":
    main()
