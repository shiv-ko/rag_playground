"""Build a lightweight diff artifact for high-confidence version pairs."""
from __future__ import annotations

import difflib
import json
import zipfile
from pathlib import Path
from typing import Any

from docx import Document as DocxDocument
from openpyxl import load_workbook
from pptx import Presentation


ROOT = Path(__file__).resolve().parent.parent
ARTIFACTS = ROOT / "artifacts"
DOCS = ROOT / "docs"
PAIRS_PATH = ARTIFACTS / "version_pairs.jsonl"
DIFF_PATH = ARTIFACTS / "version_diff_poc.jsonl"
REPORT_PATH = DOCS / "version_diff_poc.md"


def read_jsonl(path: Path) -> list[dict[str, Any]]:
    if not path.exists():
        return []
    with path.open(encoding="utf-8") as f:
        return [json.loads(line) for line in f if line.strip()]


def write_jsonl(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as f:
        for row in rows:
            f.write(json.dumps(row, ensure_ascii=False) + "\n")


def rel(path: str) -> Path:
    return ROOT / path


def clean(text: Any, max_len: int = 500) -> str:
    value = " ".join(str(text or "").split())
    if "data:image/" in value and "base64," in value:
        return "[embedded image]"
    if len(value) > max_len:
        return value[: max_len - 3] + "..."
    return value


def pptx_lines(path: Path) -> list[str]:
    prs = Presentation(path)
    lines: list[str] = []
    for slide_idx, slide in enumerate(prs.slides, start=1):
        for shape_idx, shape in enumerate(slide.shapes):
            if not getattr(shape, "has_text_frame", False):
                continue
            text = clean(shape.text)
            if text:
                lines.append(f"slide {slide_idx} shape {shape_idx}: {text}")
    return lines


def docx_lines(path: Path) -> list[str]:
    doc = DocxDocument(path)
    lines: list[str] = []
    for idx, para in enumerate(doc.paragraphs):
        text = clean(para.text)
        if text:
            lines.append(f"paragraph {idx}: {text}")
    for table_idx, table in enumerate(doc.tables):
        for row_idx, row in enumerate(table.rows):
            values = [clean(cell.text) for cell in row.cells]
            text = " | ".join(value for value in values if value)
            if text:
                lines.append(f"table {table_idx} row {row_idx}: {text}")
    return lines


def xlsx_lines(path: Path) -> list[str]:
    wb = load_workbook(path, read_only=True, data_only=True)
    lines: list[str] = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            values = [clean(cell.value) for cell in row]
            if not any(values):
                continue
            lines.append(f"{ws.title} row {row[0].row}: " + " | ".join(values))
    wb.close()
    return lines


def ipynb_lines(path: Path) -> list[str]:
    nb = json.loads(path.read_text(encoding="utf-8"))
    lines: list[str] = []
    for idx, cell in enumerate(nb.get("cells", [])):
        source = clean("".join(cell.get("source", [])))
        if source:
            lines.append(f"cell {idx} {cell.get('cell_type')}: {source}")
        for out_idx, output in enumerate(cell.get("outputs", [])):
            text = clean("".join(output.get("text", [])))
            if text:
                lines.append(f"cell {idx} output {out_idx}: {text}")
    return lines


def extract_lines(path: Path) -> list[str]:
    suffix = path.suffix.lower()
    if suffix == ".pptx":
        return pptx_lines(path)
    if suffix == ".docx":
        return docx_lines(path)
    if suffix == ".xlsx":
        return xlsx_lines(path)
    if suffix == ".ipynb":
        return ipynb_lines(path)
    return []


def diff_pair(pair: dict[str, Any]) -> dict[str, Any]:
    old_path = rel(pair["old_path"])
    new_path = rel(pair["new_path"])
    base = {
        **pair,
        "old_line_count": 0,
        "new_line_count": 0,
        "added_count": 0,
        "removed_count": 0,
        "changed_count": 0,
        "added_samples": [],
        "removed_samples": [],
        "changed_samples": [],
        "status": "ok",
        "error": None,
    }
    try:
        old_lines = extract_lines(old_path)
        new_lines = extract_lines(new_path)
    except Exception as exc:  # noqa: BLE001
        base["status"] = "failed"
        base["error"] = repr(exc)
        return base

    added: list[str] = []
    removed: list[str] = []
    changed: list[dict[str, str]] = []
    matcher = difflib.SequenceMatcher(a=old_lines, b=new_lines, autojunk=False)
    for tag, i1, i2, j1, j2 in matcher.get_opcodes():
        if tag == "equal":
            continue
        if tag == "delete":
            removed.extend(old_lines[i1:i2])
        elif tag == "insert":
            added.extend(new_lines[j1:j2])
        elif tag == "replace":
            old_chunk = old_lines[i1:i2]
            new_chunk = new_lines[j1:j2]
            for old, new in zip(old_chunk, new_chunk):
                changed.append({"before": old, "after": new})
            if len(old_chunk) > len(new_chunk):
                removed.extend(old_chunk[len(new_chunk) :])
            elif len(new_chunk) > len(old_chunk):
                added.extend(new_chunk[len(old_chunk) :])

    base.update(
        {
            "old_line_count": len(old_lines),
            "new_line_count": len(new_lines),
            "added_count": len(added),
            "removed_count": len(removed),
            "changed_count": len(changed),
            "added_samples": added[:8],
            "removed_samples": removed[:8],
            "changed_samples": changed[:8],
        }
    )
    return base


def write_report(rows: list[dict[str, Any]]) -> None:
    lines = [
        "# version_pairs 差分抽出PoC",
        "",
        "## 目的",
        "",
        "`version_pairs.jsonl` の高信頼ペアに対し、pptx/docx/xlsx/ipynbから軽量に行単位差分を抽出し、後続の意味分類と回答生成に渡せる候補を作る。",
        "",
        "## 集計",
        "",
        f"- pairs: {len(rows)}",
        f"- ok: {sum(1 for row in rows if row['status'] == 'ok')}",
        f"- failed: {sum(1 for row in rows if row['status'] != 'ok')}",
        "",
        "| project | type | old | new | changed | added | removed | status |",
        "|---|---|---|---|---:|---:|---:|---|",
    ]
    for row in rows:
        lines.append(
            f"| {row['project_name']} | {row['document_type']} | {row['old_file_name']} | {row['new_file_name']} | {row['changed_count']} | {row['added_count']} | {row['removed_count']} | {row['status']} |"
        )
    lines.extend(["", "## 差分サンプル", ""])
    for row in rows:
        lines.append(f"### {row['project_name']} / {row['document_type']} / {row['old_file_name']} -> {row['new_file_name']}")
        if row["status"] != "ok":
            lines.append(f"- failed: `{row['error']}`")
            lines.append("")
            continue
        for sample in row.get("changed_samples", [])[:3]:
            lines.append(f"- before: {sample['before']}")
            lines.append(f"  after: {sample['after']}")
        for sample in row.get("added_samples", [])[:3]:
            lines.append(f"- added: {sample}")
        for sample in row.get("removed_samples", [])[:3]:
            lines.append(f"- removed: {sample}")
        lines.append("")
    lines.extend(
        [
            "## 注意",
            "",
            "- これはPoCであり、差分の意味分類は未実装。",
            "- PPTXはshape単位、DOCXは段落/表行単位、XLSXは行単位、ipynbはセル/出力単位で比較する。",
            "- 書式だけの変更、文言だけの変更、案件遂行に関係する変更の切り分けは次段階。",
        ]
    )
    REPORT_PATH.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> None:
    pairs = read_jsonl(PAIRS_PATH)
    rows = [diff_pair(pair) for pair in pairs]
    write_jsonl(DIFF_PATH, rows)
    write_report(rows)
    print(f"pairs={len(rows)}")
    print(f"ok={sum(1 for row in rows if row['status'] == 'ok')}")
    print(f"failed={sum(1 for row in rows if row['status'] != 'ok')}")


if __name__ == "__main__":
    main()
