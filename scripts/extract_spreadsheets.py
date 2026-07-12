"""Extract structured spreadsheet artifacts for data-strategy experiments."""
from __future__ import annotations

import json
import re
import sys
import tempfile
import unicodedata
from collections import Counter
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

import openpyxl
from msoffcrypto.exceptions import InvalidKeyError
from openpyxl.cell.cell import Cell
from openpyxl.styles import Color
from openpyxl.worksheet.worksheet import Worksheet


ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from src.utils.office_crypto import (
    candidate_dates_from_filename,
    decrypt_office_file,
    derive_office_password,
    literal_password_from_filename,
    looks_like_encrypted_office_file,
)
PROJECT_ROOT = ROOT / "data" / "raw" / "share" / "共有ドライブ" / "プロジェクト"
ARTIFACTS = ROOT / "artifacts"
PROJECT_REGISTRY_PATH = ARTIFACTS / "project_registry.json"
CONTRACTS_PATH = ARTIFACTS / "contracts.jsonl"


SECTION_PREFIXES = {
    "00.": "00.提案",
    "01.": "01.契約",
    "02.": "02.計画",
    "03.": "03.データ",
    "04.": "04.分析",
    "05.": "05.会議",
    "06.": "06.報告書",
}


HEADER_KEYWORDS = (
    "タスク",
    "task",
    "担当",
    "開始",
    "終了",
    "工数",
    "ステータス",
    "進捗",
    "フェーズ",
    "マイルストーン",
    "id",
)

IGNORED_COLOR_NAMES = {None, "white", "black", "THEME:1"}


def json_default(value: Any) -> str | int | float | bool | None:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return str(value)


def write_jsonl(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as f:
        for row in rows:
            f.write(json.dumps(row, ensure_ascii=False, default=json_default) + "\n")


def color_to_hex(color: Color | None) -> str | None:
    if color is None:
        return None
    if color.type == "rgb" and color.rgb:
        rgb = str(color.rgb)
        return rgb[-6:].upper()
    if color.type == "indexed" and color.indexed is not None:
        return f"INDEXED:{color.indexed}"
    if color.type == "theme" and color.theme is not None:
        tint = "" if color.tint in (None, 0) else f":{color.tint}"
        return f"THEME:{color.theme}{tint}"
    return None


def normalize_color(hex_value: str | None) -> str | None:
    if not hex_value or not all(c in "0123456789ABCDEF" for c in hex_value) or len(hex_value) != 6:
        return hex_value
    r = int(hex_value[0:2], 16)
    g = int(hex_value[2:4], 16)
    b = int(hex_value[4:6], 16)
    if r > 220 and g > 180 and b < 120:
        return "yellow"
    if r > 180 and g < 160 and b < 160:
        return "red"
    if b > 160 and r < 180:
        return "blue"
    if r > 200 and 90 < g < 190 and b < 120:
        return "orange"
    if g > 160 and r < 180 and b < 180:
        return "green"
    if max(r, g, b) < 80:
        return "black"
    if min(r, g, b) > 220:
        return "white"
    return hex_value


def fill_color(cell: Cell) -> tuple[str | None, str | None]:
    if cell.fill is None or cell.fill.fill_type is None:
        return None, None
    raw = color_to_hex(cell.fill.fgColor)
    return raw, normalize_color(raw)


def font_color(cell: Cell) -> tuple[str | None, str | None]:
    raw = color_to_hex(cell.font.color if cell.font else None)
    return raw, normalize_color(raw)


def is_interesting_cell(cell: Cell) -> bool:
    raw_fill, _ = fill_color(cell)
    raw_font, _ = font_color(cell)
    return any(
        [
            cell.value is not None,
            raw_fill is not None,
            raw_font is not None,
            bool(cell.font and (cell.font.bold or cell.font.italic or cell.font.underline)),
        ]
    )


def get_parts(path: Path) -> tuple[str, str]:
    rel = path.relative_to(PROJECT_ROOT)
    parts = rel.parts
    project = parts[0] if len(parts) > 0 else ""
    raw_section = parts[1] if len(parts) > 1 else ""
    section = next((v for k, v in SECTION_PREFIXES.items() if raw_section.startswith(k)), raw_section)
    return project, section


def base_metadata(path: Path) -> dict[str, Any]:
    project, section = get_parts(path)
    return {
        "source_path": str(path.relative_to(ROOT)),
        "project_name": project,
        "section": section,
        "file_name": path.name,
    }


def sheet_dimensions(ws: Worksheet) -> dict[str, Any]:
    return {
        "sheet_name": ws.title,
        "max_row": ws.max_row,
        "max_column": ws.max_column,
        "auto_filter_ref": ws.auto_filter.ref,
        "freeze_panes": str(ws.freeze_panes) if ws.freeze_panes else None,
        "merged_ranges": [str(rng) for rng in ws.merged_cells.ranges],
        "hidden_rows": [idx for idx, dim in ws.row_dimensions.items() if dim.hidden],
        "hidden_columns": [idx for idx, dim in ws.column_dimensions.items() if dim.hidden],
    }


def normalize_project_text(text: str) -> str:
    """build_contract_registry.pyのnormalize_textと同一ロジック（NFC正規化＋全角スペース→半角）。

    scripts/build_*.py群は意図的に自己完結スタイルなのでcross-script importはせず複製する。
    """
    return unicodedata.normalize("NFC", text).replace("　", " ")


def load_primary_aliases(project_registry_path: Path = PROJECT_REGISTRY_PATH) -> dict[str, str]:
    """project_registry.jsonから`project_name(正規化済み) -> primary_alias`の対応表を作る。"""
    if not project_registry_path.exists():
        return {}
    data = json.loads(project_registry_path.read_text(encoding="utf-8"))
    return {
        normalize_project_text(row["project_name"]): row["primary_alias"]
        for row in data
        if row.get("primary_alias")
    }


def candidate_start_dates_from_contracts(
    project_name: str, contracts_path: Path = CONTRACTS_PATH
) -> list[str]:
    """同一案件のcontracts.jsonlから開始日候補(YYYYMMDD)を抽出する。

    スケジュールxlsx自身の復号に使うため、schedule_tasks.jsonl由来の日付候補
    （build_contract_registry.pyのcandidate_dates_from_schedule）は循環参照になり使えない。
    かえで案件の契約書は既に復号済み・contracts.jsonlにstart_dateが載っているため、
    こちらを日付候補源にする。
    """
    if not contracts_path.exists():
        return []
    normalized_project = normalize_project_text(project_name)
    candidates: list[str] = []
    seen: set[str] = set()
    with contracts_path.open(encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            row = json.loads(line)
            if normalize_project_text(row.get("project_name", "")) != normalized_project:
                continue
            for key in ("start_date", "end_date"):
                value = row.get(key)
                if not value:
                    continue
                digits = re.sub(r"[^0-9]", "", str(value))[:8]
                if len(digits) == 8 and digits not in seen:
                    seen.add(digits)
                    candidates.append(digits)
    return candidates


def attempt_decrypt_workbook(
    path: Path,
    project_name: str,
    primary_aliases: dict[str, str],
    output_path: Path,
    contracts_path: Path = CONTRACTS_PATH,
) -> bool:
    """暗号化されたxlsxに対し、候補パスワードを順に試して復号し、output_pathに平文コピーを書く。

    build_contract_registry.pyのattempt_decrypt_contract_textと同じ2系統の候補:
    (1) ファイル名`pw-<トークン>`のトークン自体をリテラルパスワードとして試す。
    (2) DA-規則`DA-[案件略号]-[開始年月日8桁]-[拡張子]`。開始年月日はファイル名の
        `pw-...<8桁>`命名慣習、無ければ同一案件のcontracts.jsonl（既に復号済みの契約書）
        のstart_date/end_dateを候補にする。
    全候補が`InvalidKeyError`（パスワード誤り）で失敗した場合はFalseを返す。
    """
    literal_password = literal_password_from_filename(path)
    if literal_password is not None:
        try:
            decrypt_office_file(path, literal_password, output_path)
            return True
        except InvalidKeyError:
            pass

    alias = primary_aliases.get(normalize_project_text(project_name))
    if not alias:
        return False
    candidates = candidate_dates_from_filename(path)
    if not candidates:
        candidates = candidate_start_dates_from_contracts(project_name, contracts_path)
    ext = path.suffix
    for candidate in candidates:
        password = derive_office_password(alias, candidate, ext)
        try:
            decrypt_office_file(path, password, output_path)
            return True
        except InvalidKeyError:
            continue
    return False


def load_workbook(
    path: Path,
    primary_aliases: dict[str, str] | None = None,
    contracts_path: Path = CONTRACTS_PATH,
):
    """通常どおりworkbookを開く。暗号化(CDFV2)らしい失敗の場合のみ復号を試みて再ロードする。

    非暗号化の破損等、他の原因での失敗は復号を試みず元の例外をそのまま送出する
    （main()側の既存failures収集ロジックに影響しない）。
    """
    try:
        return openpyxl.load_workbook(
            path,
            data_only=True,
            read_only=False,
            keep_vba=False,
            keep_links=False,
        )
    except Exception:
        if primary_aliases is None or not looks_like_encrypted_office_file(path):
            raise
        project_name, _ = get_parts(path)
        with tempfile.TemporaryDirectory() as tmp_dir:
            decrypted_path = Path(tmp_dir) / f"decrypted{path.suffix}"
            if not attempt_decrypt_workbook(
                path, project_name, primary_aliases, decrypted_path, contracts_path
            ):
                raise
            return openpyxl.load_workbook(
                decrypted_path,
                data_only=True,
                read_only=False,
                keep_vba=False,
                keep_links=False,
            )


def likely_header_row(ws: Worksheet) -> int | None:
    best_row = None
    best_score = 0
    max_scan = min(ws.max_row, 30)
    for row_idx in range(1, max_scan + 1):
        values = [str(cell.value).strip().lower() for cell in ws[row_idx] if cell.value is not None]
        score = sum(any(keyword in value for keyword in HEADER_KEYWORDS) for value in values)
        if score > best_score:
            best_score = score
            best_row = row_idx
    return best_row if best_score >= 2 else None


def extract_schedule_rows(path: Path, ws: Worksheet, meta: dict[str, Any]) -> list[dict[str, Any]]:
    if meta["section"] != "02.計画":
        return []
    header_row = likely_header_row(ws)
    if header_row is None:
        return []

    headers: list[str | None] = []
    for cell in ws[header_row]:
        value = str(cell.value).strip() if cell.value is not None else ""
        headers.append(value or None)

    rows: list[dict[str, Any]] = []
    carry_values: dict[str, Any] = {}
    carry_headers = ("フェーズNo.", "フェーズNo", "フェーズ名", "マイルストーンID", "MS", "工程", "カテゴリ")
    for row_idx in range(header_row + 1, ws.max_row + 1):
        values: dict[str, Any] = {}
        row_fill_colors = []
        has_value = False
        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=row_idx, column=col_idx)
            fill_raw, fill_name = fill_color(cell)
            if fill_name not in IGNORED_COLOR_NAMES:
                row_fill_colors.append(fill_name)
            if header:
                value = cell.value
                if value is None and header in carry_values:
                    value = carry_values[header]
                elif value is not None and header in carry_headers:
                    carry_values[header] = value
                values[header] = value
            if cell.value is not None:
                has_value = True
        if not has_value:
            continue
        rows.append(
            {
                **meta,
                "sheet_name": ws.title,
                "header_row": header_row,
                "row_number": row_idx,
                "row_fill_colors": sorted(set(row_fill_colors)),
                "dominant_row_fill": Counter(row_fill_colors).most_common(1)[0][0] if row_fill_colors else None,
                "values": values,
            }
        )
    return rows


def extract_workbook(
    path: Path,
    primary_aliases: dict[str, str] | None = None,
    contracts_path: Path = CONTRACTS_PATH,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    meta = base_metadata(path)
    value_wb = load_workbook(path, primary_aliases, contracts_path)
    sheets: list[dict[str, Any]] = []
    cells: list[dict[str, Any]] = []
    highlights: list[dict[str, Any]] = []
    schedule_rows: list[dict[str, Any]] = []

    for sheet_name in value_wb.sheetnames:
        ws = value_wb[sheet_name]
        sheets.append({**meta, **sheet_dimensions(ws)})
        schedule_rows.extend(extract_schedule_rows(path, ws, meta))

        for row in ws.iter_rows():
            for cell in row:
                if not is_interesting_cell(cell):
                    continue
                fill_raw, fill_name = fill_color(cell)
                font_raw, font_name = font_color(cell)
                row_data = {
                    **meta,
                    "sheet_name": sheet_name,
                    "cell": cell.coordinate,
                    "row": cell.row,
                    "column": cell.column,
                    "value": cell.value,
                    "formula": None,
                    "number_format": cell.number_format,
                    "fill_color": fill_raw,
                    "fill_color_name": fill_name,
                    "font_color": font_raw,
                    "font_color_name": font_name,
                    "bold": bool(cell.font and cell.font.bold),
                    "italic": bool(cell.font and cell.font.italic),
                    "underline": bool(cell.font and cell.font.underline),
                }
                cells.append(row_data)
                has_mark = (
                    fill_name not in IGNORED_COLOR_NAMES
                    or font_name not in IGNORED_COLOR_NAMES
                    or row_data["bold"]
                    or row_data["italic"]
                    or row_data["underline"]
                )
                if has_mark:
                    highlights.append(row_data)

    value_wb.close()
    return sheets, cells, highlights, schedule_rows


def main() -> None:
    xlsx_files = sorted(
        p
        for p in PROJECT_ROOT.rglob("*.xlsx")
        if not p.name.startswith("~$") and "__pycache__" not in p.parts
        and "02.計画" in p.parts
    )
    all_sheets: list[dict[str, Any]] = []
    all_cells: list[dict[str, Any]] = []
    all_highlights: list[dict[str, Any]] = []
    all_schedule_rows: list[dict[str, Any]] = []
    failures: list[dict[str, str]] = []
    primary_aliases = load_primary_aliases()

    for path in xlsx_files:
        print(f"extracting {path.relative_to(ROOT)}", flush=True)
        try:
            sheets, cells, highlights, schedule_rows = extract_workbook(path, primary_aliases)
        except Exception as exc:  # noqa: BLE001 - keep extraction audit complete.
            failures.append({"source_path": str(path.relative_to(ROOT)), "error": repr(exc)})
            continue
        all_sheets.extend(sheets)
        all_cells.extend(cells)
        all_highlights.extend(highlights)
        all_schedule_rows.extend(schedule_rows)

    write_jsonl(ARTIFACTS / "spreadsheet_sheets.jsonl", all_sheets)
    write_jsonl(ARTIFACTS / "spreadsheet_cells.jsonl", all_cells)
    write_jsonl(ARTIFACTS / "highlight_cells.jsonl", all_highlights)
    write_jsonl(ARTIFACTS / "schedule_tasks.jsonl", all_schedule_rows)
    write_jsonl(ARTIFACTS / "spreadsheet_failures.jsonl", failures)

    print(f"xlsx_files={len(xlsx_files)}")
    print(f"sheets={len(all_sheets)}")
    print(f"cells={len(all_cells)}")
    print(f"highlight_cells={len(all_highlights)}")
    print(f"schedule_rows={len(all_schedule_rows)}")
    print(f"failures={len(failures)}")


if __name__ == "__main__":
    main()
