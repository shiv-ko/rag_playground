"""Build contract registry from project contract documents."""
from __future__ import annotations

import json
import re
import sys
import tempfile
import unicodedata
from datetime import date, timedelta
from pathlib import Path
from typing import Any

from docx import Document as DocxDocument
from msoffcrypto.exceptions import InvalidKeyError
from pptx import Presentation

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from src.utils.office_crypto import (
    candidate_dates_from_filename,
    decrypt_office_file,
    derive_office_password,
    literal_password_from_filename,
    load_primary_aliases,
    looks_like_encrypted_office_file,
)

SHARE_ROOT = ROOT / "data" / "raw" / "share" / "共有ドライブ"
PROJECT_ROOT = SHARE_ROOT / "プロジェクト"
ARTIFACTS = ROOT / "artifacts"
CONTRACTS_PATH = ARTIFACTS / "contracts.jsonl"
PROJECT_REGISTRY_PATH = ARTIFACTS / "project_registry.json"
SCHEDULE_TASKS_PATH = ARTIFACTS / "schedule_tasks.jsonl"

# looks_like_encrypted_office_file / candidate_dates_from_filename / literal_password_from_filename
# は src/utils/office_crypto.py へ移動済み（extract_spreadsheets.pyとの共有のため）。
# この行より下で使用する名前はそちらからのimportで解決される。


def normalize_text(text: str) -> str:
    return unicodedata.normalize("NFC", text).replace("\u3000", " ")


def parse_yen(value: str) -> int:
    return int(re.sub(r"[^0-9]", "", value))


def parse_float(value: str) -> float:
    return float(value.replace(",", ""))


def read_docx_text(path: Path) -> str:
    doc = DocxDocument(path)
    lines = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if values:
                lines.append(" | ".join(values))
    return normalize_text("\n".join(lines))


def read_pptx_text(path: Path) -> str:
    prs = Presentation(path)
    lines: list[str] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text.strip():
                lines.append(shape.text)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    values = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if values:
                        lines.append(" | ".join(values))
    return normalize_text("\n".join(lines))


def read_pdf_text(path: Path) -> str:
    import pypdf

    reader = pypdf.PdfReader(str(path))
    return normalize_text("\n".join(page.extract_text() or "" for page in reader.pages))


# load_primary_aliases はextract_spreadsheets.pyと独立実装が分岐していたため
# src/utils/office_crypto.py へ統合済み（DA規則やレジストリ仕様の変更時に片方だけ直る
# 事故を防ぐ）。normalize_text（NFC正規化＋全角スペース→半角）と同一ロジックを内部で使う
# ため、この関数の返り値・挙動は従来のローカル実装と変わらない。


def candidate_dates_from_schedule(
    project_name: str, schedule_tasks_path: Path = SCHEDULE_TASKS_PATH
) -> list[str]:
    """同一案件のスケジュールregistry(schedule_tasks.jsonl)から日付候補(YYYYMMDD)を抽出する。

    暗号化された契約書自身からは開始日を読めない（鶏と卵）ため、ファイル名に日付候補が無い場合の
    フォールバック候補源として、案件ごとに既に構造化済みの他registryの日付を順に試す。
    """
    if not schedule_tasks_path.exists():
        return []
    normalized_project = normalize_text(project_name)
    candidates: list[str] = []
    seen: set[str] = set()
    with schedule_tasks_path.open(encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            row = json.loads(line)
            if normalize_text(row.get("project_name", "")) != normalized_project:
                continue
            values = row.get("values") or {}
            for key in ("開始日", "終了日"):
                value = values.get(key)
                if not value:
                    continue
                digits = re.sub(r"[^0-9]", "", str(value))[:8]
                if len(digits) == 8 and digits not in seen:
                    seen.add(digits)
                    candidates.append(digits)
    return candidates


def _try_decrypt_with_password(path: Path, password: str) -> str | None:
    """1つのパスワード候補で復号を試み、成功すれば平文テキストを返す（誤りならNone）。"""
    ext = path.suffix
    with tempfile.TemporaryDirectory() as tmp_dir:
        output_path = Path(tmp_dir) / f"decrypted{ext}"
        try:
            decrypt_office_file(path, password, output_path)
        except InvalidKeyError:
            return None
        return read_docx_text(output_path)


def attempt_decrypt_contract_text(
    path: Path,
    project_name: str,
    primary_aliases: dict[str, str],
    schedule_tasks_path: Path = SCHEDULE_TASKS_PATH,
) -> str | None:
    """暗号化された契約書docxに対し、候補パスワードを順に試して復号し、平文テキストを返す。

    実データで2種類の命名慣習が確認されている:
    (1) `pw-<トークン>`のトークン自体がDA-規則を介さずそのまま平文パスワードになっているケース
        （かえで案件の契約書で確認済み）。まずこちらを試す（案件略号・日付を必要としない、
        最も直接的な候補）。
    (2) 規定の `DA-[案件略号]-[開始年月日8桁]-[拡張子]` 形式。開始年月日はそのファイル自身の
        契約期間フィールドにしか書かれておらず、暗号化されたファイル自身からは読めない
        （鶏と卵）ため、(2-a)ファイル名の`pw-...<8桁>`命名慣習から日付候補を抽出し、
        それが見つからない場合のみ(2-b)同一案件のスケジュールregistry
        (schedule_tasks.jsonl)の開始日/終了日候補を順に試す。
    全候補が`InvalidKeyError`（パスワード誤り）で失敗した場合はNoneを返し、
    呼び出し側は既存どおり安全側のstatus=failedにフォールバックする。
    """
    literal_password = literal_password_from_filename(path)
    if literal_password is not None:
        text = _try_decrypt_with_password(path, literal_password)
        if text is not None:
            return text

    alias = primary_aliases.get(normalize_text(project_name))
    if not alias:
        return None
    candidates = candidate_dates_from_filename(path)
    if not candidates:
        candidates = candidate_dates_from_schedule(project_name, schedule_tasks_path)
    ext = path.suffix
    for candidate in candidates:
        password = derive_office_password(alias, candidate, ext)
        text = _try_decrypt_with_password(path, password)
        if text is not None:
            return text
    return None


def write_jsonl(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as f:
        for row in rows:
            f.write(json.dumps(row, ensure_ascii=False) + "\n")


def select_contract_path(project_dir: Path) -> Path | None:
    contract_dirs = [p for p in project_dir.iterdir() if p.is_dir() and p.name.startswith("01.")]
    candidates: list[Path] = []
    for contract_dir in contract_dirs:
        candidates.extend(
            p for p in contract_dir.glob("*.docx")
            if not p.name.startswith("~$") and "draft" not in p.stem.lower()
        )
    if not candidates:
        return None
    return sorted(candidates, key=lambda p: ("pw-" in p.name, p.name))[0]


def extract_billing_clause(text: str) -> str:
    match = re.search(r"6\.\s*報酬および支払条件(?P<body>.*?)(?:\n7\.|\n第7条|$)", text, re.DOTALL)
    body = match.group("body") if match else text
    return " ".join(body.split())[:1200]


def classify_contract_type(text: str) -> str:
    if "time_and_materials" in text or "Time & Materials" in text:
        return "time_and_materials"
    if "固定価格" in text or "固定金額" in text or "金額を固定" in text:
        return "fixed"
    return "unknown"


def classify_rounding_rule(clause: str) -> str:
    if "30分未満を0.5時間" in clause and "30分超60分未満を1.0時間" in clause:
        return "half_unit_bucket"
    if "30分" in clause and ("切り上げ" in clause or "切上げ" in clause):
        return "ceiling_to_unit"
    if "丸め" not in clause and "端数" not in clause:
        return "none"
    return "unknown"


def first_amount(patterns: list[str], text: str) -> int | None:
    for pattern in patterns:
        match = re.search(pattern, text)
        if match:
            return parse_yen(match.group(1))
    return None


def extract_advance_payment(text: str, amount_incl: int | None) -> int | None:
    direct = first_amount([r"着手金[：:\s]*([0-9,]+)円"], text)
    if direct is not None:
        return direct
    if amount_incl is None:
        return None
    # 支払表は案件ごとに列順が異なる(税抜/税込の位置が揺れる)ため、金額セルを直接
    # 拾わず、同じ行にある比率(%)から契約金額（税込）に対する割合として算出する。
    for line in text.splitlines():
        if "着手金" not in line:
            continue
        match = re.search(r"([0-9]+(?:\.[0-9]+)?)\s*[%％]", line)
        if match:
            return int(round(amount_incl * float(match.group(1)) / 100))
    return None


def extract_dates(text: str) -> tuple[str | None, str | None, int | None]:
    match = re.search(r"(\d{4})[-年/](\d{1,2})[-月/](\d{1,2})日?\s*から\s*(\d{4})[-年/](\d{1,2})[-月/](\d{1,2})日?\s*まで", text)
    if match:
        values = [int(v) for v in match.groups()]
        start = date(values[0], values[1], values[2])
        end = date(values[3], values[4], values[5])
        return start.isoformat(), end.isoformat(), (end - start).days + 1
    match = re.search(r"(\d{4})[-年/](\d{1,2})[-月/](\d{1,2})日?\s*から起算して(\d+)週間", text)
    if not match:
        return None, None, None
    year, month, day, weeks = [int(v) for v in match.groups()]
    start = date(year, month, day)
    end = start + timedelta(days=weeks * 7 - 1)
    return start.isoformat(), end.isoformat(), (end - start).days + 1


def extract_report_values(project_dir: Path) -> dict[str, Any]:
    report_dirs = [p for p in project_dir.iterdir() if p.is_dir() and p.name.startswith("06.")]
    texts: list[str] = []
    # 「完了案件」判定用のファイル存在フラグ。final_amount_incl_tax（金額の正規表現抽出）とは
    # 独立に、06.報告書配下にoldを除く最終報告ファイルが1件でも存在すれば真とする。
    # 固定価格契約は最終報告書で金額を再掲しない/言い回しが違うため、final_amount_incl_taxの
    # 非null性だけでは「完了」の代理指標として不十分（実データで確認済みのバグ）。
    has_final_report = False
    for report_dir in report_dirs:
        for path in sorted(report_dir.glob("*.pptx")):
            if "old" in path.stem.lower() or path.name.startswith("~$"):
                continue
            has_final_report = True
            try:
                texts.append(read_pptx_text(path))
            except Exception:
                continue
        for path in sorted(report_dir.glob("*.pdf")):
            if "old" not in path.stem.lower():
                has_final_report = True
            try:
                texts.append(read_pdf_text(path))
            except Exception:
                continue
    text = "\n".join(texts)
    actual_hours = None
    m_hours = re.search(r"実績工数[：:\s|]*([0-9]+(?:\.[0-9]+)?)\s*時間", text)
    if m_hours:
        actual_hours = parse_float(m_hours.group(1))
    final_incl = first_amount([r"最終請求金額（税込）[：:\s|]*([0-9,]+)\s*円", r"税込金額[：:\s|]*([0-9,]+)\s*円"], text)
    return {
        "actual_hours": actual_hours,
        "final_amount_incl_tax": final_incl,
        "has_final_report": has_final_report,
    }


def extract_payment_schedule(text: str) -> list[dict[str, Any]]:
    """支払回テーブルから{month, amount_incl_tax, due_date}のリストを抽出する（Q40用）。

    列名・列順は案件ごとに揺れる（支払回/名目/比率/... のバリエーションが複数確認済み）ため、
    ヘッダー行を都度探索し「税込」を含む列・「支払期」を含む列（支払期日/支払期限どちらも
    カバー）をインデックスで解決する。支払期日列に日付が無く条件文と複合値になっているケース
    （例:「...5営業日以内（2025-09-24）」）にも対応するため、セル内のYYYY-MM-DD部分文字列を
    正規表現で探す。
    """
    lines = text.splitlines()
    header_idx = None
    for i, line in enumerate(lines):
        if "|" in line and "支払回" in line:
            header_idx = i
            break
    if header_idx is None:
        return []
    headers = [h.strip() for h in lines[header_idx].split("|")]
    amount_idx = next((i for i, h in enumerate(headers) if "税込" in h), None)
    due_idx = next((i for i, h in enumerate(headers) if "支払期" in h), None)
    if amount_idx is None or due_idx is None:
        return []
    schedule: list[dict[str, Any]] = []
    for line in lines[header_idx + 1:]:
        if "|" not in line:
            break
        cells = [c.strip() for c in line.split("|")]
        if len(cells) != len(headers):
            break
        amount_match = re.search(r"([0-9,]+)\s*円", cells[amount_idx])
        date_match = re.search(r"(\d{4}-\d{2}-\d{2})", cells[due_idx])
        if not amount_match or not date_match:
            continue
        due_date = date_match.group(1)
        schedule.append({
            "month": due_date[:7],
            "amount_incl_tax": parse_yen(amount_match.group(1)),
            "due_date": due_date,
        })
    return schedule


def parse_contract_text(project_name: str, source_path: str, text: str) -> dict[str, Any]:
    text = normalize_text(text)
    clause = extract_billing_clause(text)
    contract_type = classify_contract_type(text)
    start_date, end_date, days = extract_dates(text)
    rate = first_amount([r"時間単価[は：:\s|]*([0-9,]+)円", r"1時間当たり([0-9,]+)円"], text)
    esth_match = re.search(r"(?:想定総工数|見込工数)[は：:\s]*([0-9]+(?:\.[0-9]+)?)\s*時間", text)
    esth = parse_float(esth_match.group(1)) if esth_match else None
    amount_excl = first_amount(
        [
            r"(?:契約金額|報酬総額|見込金額|想定金額)（税抜）[：:\s|]*([0-9,]+)円",
            r"税抜(?:契約金額|報酬総額|見込金額|想定金額)[：:\s|]*([0-9,]+)円",
            r"税抜\s*([0-9,]+)円",
        ],
        text,
    )
    amount_incl = first_amount(
        [
            r"(?:契約金額|報酬総額|見込金額|想定金額)（税込）[：:\s|]*([0-9,]+)円",
            r"税込(?:契約金額|報酬総額|見込金額|想定金額)[：:\s|]*([0-9,]+)円",
            r"税込\s*([0-9,]+)円",
        ],
        text,
    )
    tax = first_amount([r"消費税(?:額)?[：:\s]*([0-9,]+)円"], text)
    advance = extract_advance_payment(text, amount_incl)
    return {
        "project_name": project_name,
        "contract_type": contract_type,
        "rate_yen_per_hour": rate,
        "esth_hours": esth,
        "estimated_amount_excl_tax": amount_excl,
        "estimated_amount_incl_tax": amount_incl,
        "tax_rate": round(tax / amount_excl, 2) if tax and amount_excl else 0.10,
        "rounding_unit_minutes": 30 if "30分" in clause else None,
        "rounding_rule": classify_rounding_rule(clause),
        "start_date": start_date,
        "end_date": end_date,
        "contract_period_days": days,
        "advance_payment_amount": advance,
        "payment_schedule": extract_payment_schedule(text),
        "raw_billing_clause": clause,
        "source_path": source_path,
        "status": "ok",
        "error": None,
    }


def build_contract_registry(
    project_root: Path = PROJECT_ROOT,
    project_registry_path: Path = PROJECT_REGISTRY_PATH,
    schedule_tasks_path: Path = SCHEDULE_TASKS_PATH,
) -> list[dict[str, Any]]:
    primary_aliases = load_primary_aliases(project_registry_path)
    rows: list[dict[str, Any]] = []
    for project_dir in sorted(p for p in project_root.iterdir() if p.is_dir()):
        path = select_contract_path(project_dir)
        if path is None:
            continue
        project_name = normalize_text(project_dir.name)
        try:
            rel_path = str(path.relative_to(ROOT))
        except ValueError:
            rel_path = str(path)
        try:
            row = parse_contract_text(project_name, rel_path, read_docx_text(path))
            row.update(extract_report_values(project_dir))
        except Exception as exc:  # noqa: BLE001
            decrypted_text = None
            if looks_like_encrypted_office_file(path):
                try:
                    decrypted_text = attempt_decrypt_contract_text(
                        path, project_name, primary_aliases, schedule_tasks_path
                    )
                except Exception:  # noqa: BLE001
                    decrypted_text = None
            if decrypted_text is not None:
                try:
                    row = parse_contract_text(project_name, rel_path, decrypted_text)
                    row.update(extract_report_values(project_dir))
                except Exception as inner_exc:  # noqa: BLE001
                    row = {
                        "project_name": project_name,
                        "source_path": rel_path,
                        "contract_type": "unknown",
                        "status": "failed",
                        "error": repr(inner_exc),
                    }
            else:
                row = {
                    "project_name": project_name,
                    "source_path": rel_path,
                    "contract_type": "unknown",
                    "status": "failed",
                    "error": repr(exc),
                }
        rows.append(row)
    return rows


def main() -> None:
    rows = build_contract_registry()
    write_jsonl(CONTRACTS_PATH, rows)
    print(f"contracts={len(rows)}")
    print(f"ok={sum(1 for row in rows if row.get('status') == 'ok')}")
    print(f"failed={sum(1 for row in rows if row.get('status') != 'ok')}")


if __name__ == "__main__":
    main()
