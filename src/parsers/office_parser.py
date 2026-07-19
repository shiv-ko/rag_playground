"""Word/Excel/PowerPoint パーサー。各ライブラリが未インストールならスタブを返す。"""
import logging
import tempfile
from pathlib import Path

from src.models import Document
from src.utils.office_crypto import (
    decrypt_office_file,
    load_primary_aliases,
    looks_like_encrypted_office_file,
    password_candidates_for_file,
    project_name_from_path,
)

logger = logging.getLogger(__name__)

SUPPORTED = {".docx", ".xlsx", ".pptx"}

# 実データ本体を直接パースする本番の索引経路（ParserDispatcher経由）向けの
# 暗号化ファイル対応。scripts/extract_spreadsheets.py・scripts/build_contract_registry.py
# は別経路（構造化artifact生成専用）で既にCDFV2内容ベース検知＋DA規則復号を持つが、
# 実際の検索インデックスを作るこちらの経路には一切配線されておらず、KAEDEのスケジュール.xlsx
# のようにファイル名に`pw-`マーカーが無い暗号化ファイルはチャンクが1件も生成されない
# まま欠落していた（docs/test_missing_text_only_diag_20260718.md 修正候補2）。
# artifacts/はrun_pipeline側で事前生成される機械生成レジストリの読み込み専用として使う
# （このモジュールから再生成はしない）。
_ROOT = Path(__file__).resolve().parents[2]
PROJECT_REGISTRY_PATH = _ROOT / "artifacts" / "project_registry.json"
CONTRACTS_PATH = _ROOT / "artifacts" / "contracts.jsonl"


class OfficeParser:
    def can_handle(self, file_path: Path) -> bool:
        return file_path.suffix.lower() in SUPPORTED

    def _decrypt_via_password_candidates(self, file_path: Path, output_path: Path) -> bool:
        """CDFV2(OLE)コンテナ署名で暗号化らしいと判定できる場合のみ、案件名・案件略号・
        contracts.jsonlの開始/終了日から導出したパスワード候補を順に試して復号する。

        ファイル名に`pw-`マーカーが無くても、ディレクトリ構造上の案件名
        （`.../プロジェクト/<案件名>/...`）とartifacts/project_registry.json・
        artifacts/contracts.jsonl（機械生成の読み込み専用レジストリ）だけから
        候補を導出する汎用ロジック。特定の案件名・ファイル名へのハードコードはしない。
        署名が無い（＝暗号化以外の破損原因）場合や、候補が尽きても復号できない場合は
        Falseを返し、呼び出し側の既存stubフォールバックに委ねる。

        `InvalidKeyError`（パスワード誤り）だけでなく、その親クラス`DecryptionError`
        （暗号化情報ストリームの解釈失敗等、他の理由での単一候補の復号失敗）も候補単位で
        握りつぶし次の候補へ進む。scripts/extract_spreadsheets.py・
        scripts/build_contract_registry.pyの既存箇所は候補源が限定的（xlsx/docxそれぞれ
        単一形式）で`InvalidKeyError`のみ捕捉する運用だが、こちらは複数形式・複数案件を
        横断する汎用経路のため、1候補の想定外の失敗で復号フォールバック全体を諦めない
        よう広めに捕捉する。
        """
        from msoffcrypto.exceptions import DecryptionError

        if not looks_like_encrypted_office_file(file_path):
            return False

        primary_aliases = load_primary_aliases(PROJECT_REGISTRY_PATH)
        project_name = project_name_from_path(file_path) or ""
        candidates = password_candidates_for_file(
            file_path, project_name, primary_aliases, CONTRACTS_PATH
        )
        for password in candidates:
            try:
                decrypt_office_file(file_path, password, output_path)
                return True
            except DecryptionError:
                continue
        return False

    def _load_docx(self, file_path: Path, docx):
        """通常どおりdocxを開く。CDFV2署名で暗号化らしいと判定できる失敗のときだけ
        復号を試みて再ロードする。それ以外の原因（単純な破損等）は元の例外をそのまま送出し、
        呼び出し側の既存stubフォールバックに委ねる（動作を変えない）。"""
        try:
            return docx.Document(str(file_path))
        except Exception:
            with tempfile.TemporaryDirectory() as tmp_dir:
                decrypted_path = Path(tmp_dir) / f"decrypted{file_path.suffix}"
                if not self._decrypt_via_password_candidates(file_path, decrypted_path):
                    raise
                return docx.Document(str(decrypted_path))

    def parse(self, file_path: Path) -> list[Document]:
        suffix = file_path.suffix.lower()
        if suffix == ".docx":
            return self._parse_docx(file_path)
        if suffix == ".xlsx":
            return self._parse_xlsx(file_path)
        if suffix == ".pptx":
            return self._parse_pptx(file_path)
        return []

    def _parse_docx(self, file_path: Path) -> list[Document]:
        try:
            import docx
        except ImportError:
            return [Document(text=f"[DOCX未解析: python-docx未インストール] {file_path.name}",
                             source_path=file_path, location="stub")]
        try:
            from docx.table import Table
            from docx.text.paragraph import Paragraph

            doc = self._load_docx(file_path, docx)
            texts: list[str] = []
            # iter_inner_content()は本文中の段落/表の出現順（body順）をそのまま
            # 返す公開API。生XML要素を手動でqn判定・再構築する必要がない。
            for item in doc.iter_inner_content():
                try:
                    if isinstance(item, Table):
                        texts.extend(self._extract_docx_table_rows(item, file_path))
                    elif isinstance(item, Paragraph):
                        if item.text.strip():
                            texts.append(item.text)
                except Exception:
                    # 個別の段落/表の解釈に失敗しても文書全体は読み進める
                    logger.warning(
                        "docx本文要素(段落/表)の抽出に失敗しスキップ: %s", file_path,
                        exc_info=True,
                    )
                    continue
            text = "\n".join(texts)
        except Exception:
            logger.warning("docx全体の解析に失敗: %s", file_path, exc_info=True)
            return [Document(
                text=f"[DOCX解析失敗: パスワード保護または破損の可能性] {file_path.name}",
                source_path=file_path,
                location="stub",
            )]
        return [Document(text=text, source_path=file_path, location="full")]

    def _extract_docx_table_rows(self, table, file_path: Path) -> list[str]:
        """Word表を行ごとに ' | ' 結合した文字列のリストへ変換する。

        結合セル・空セル・ネストした表があっても例外を出さないよう、各要素の
        アクセスを個別に防御し、セルの中身は常に文字列化して扱う。

        python-docxの`row.cells`は結合セルを実XML要素(tc)の同一性ではなく
        グリッド位置ベースで返すため、横結合(gridSpan)はスパン列数分、
        縦結合(vMerge)は継続行の分だけ同一tc要素を複数回yieldする。
        そのまま連結すると同じ内容が重複するため、tc要素で表単位にデデュープ
        する（表ごとに独立したスコープ。ネスト表は別スコープ）。

        注意: `id(cell._tc)`のようにint化して集合に入れてはいけない。
        lxmlのElementプロキシは短命な使い捨てオブジェクトとして再生成される
        ことがあり、参照を保持しないとGC直後にCPythonのアロケータが同じ
        メモリアドレスを別の（無関係な）オブジェクトに再利用するため、
        全く別のセル同士がid衝突で誤って「重複」判定されてしまう
        （実データの社内用語集.docxで実証: 123セル中119セルが誤って重複扱い
        されデータの大半が消失した）。そのためtc要素オブジェクト自体を
        集合に保持し、GCによるアドレス再利用を防ぐ。
        """
        rows_text: list[str] = []
        seen_tcs: set = set()
        try:
            rows = table.rows
        except Exception:
            logger.warning("docx表の行取得に失敗しスキップ: %s", file_path, exc_info=True)
            return rows_text
        for row in rows:
            try:
                row_cells = row.cells
            except Exception:
                logger.warning(
                    "docx表の行のセル取得に失敗しスキップ: %s", file_path, exc_info=True
                )
                continue
            cell_parts: list[str] = []
            nested_parts: list[str] = []
            for cell in row_cells:
                try:
                    tc = cell._tc  # noqa: SLF001 - 結合セルの同一性判定に必要
                except Exception:
                    tc = None
                if tc is not None:
                    if tc in seen_tcs:
                        # 横結合/縦結合による同一tc要素の重複yieldをスキップ
                        continue
                    seen_tcs.add(tc)
                try:
                    cell_text = " ".join(str(cell.text).split())
                except Exception:
                    logger.warning(
                        "docx表セルのテキスト取得に失敗しスキップ: %s", file_path,
                        exc_info=True,
                    )
                    cell_text = ""
                if cell_text:
                    cell_parts.append(cell_text)
                try:
                    nested_tables = cell.tables
                except Exception:
                    logger.warning(
                        "docx表セルのネスト表取得に失敗しスキップ: %s", file_path,
                        exc_info=True,
                    )
                    nested_tables = []
                for nested in nested_tables:
                    nested_parts.extend(self._extract_docx_table_rows(nested, file_path))
            if cell_parts:
                rows_text.append(" | ".join(cell_parts))
            rows_text.extend(nested_parts)
        return rows_text

    def _extract_xlsx_docs(self, wb, file_path: Path) -> list[Document]:
        docs = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = []
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) if c is not None else "" for c in row]
                if any(c.strip() for c in cells):
                    rows.append(" | ".join(cells))
            if rows:
                docs.append(Document(
                    text="\n".join(rows),
                    source_path=file_path,
                    location=f"sheet_{sheet_name}",
                ))
        return docs or [Document(text="", source_path=file_path, location="empty")]

    def _parse_xlsx(self, file_path: Path) -> list[Document]:
        try:
            import openpyxl
        except ImportError:
            return [Document(text=f"[XLSX未解析: openpyxl未インストール] {file_path.name}",
                             source_path=file_path, location="stub")]
        try:
            wb = openpyxl.load_workbook(str(file_path), read_only=True, data_only=True)
            return self._extract_xlsx_docs(wb, file_path)
        except Exception:
            pass
        # 通常ロードが失敗した場合のみ、CDFV2署名ベースで暗号化を疑い復号を試みる。
        # ワークブック処理までtempdirのスコープ内で完結させる（read_only=Trueのworkbookは
        # 遅延読み込みのため、復号済み一時ファイルが削除される前に処理を終える必要がある）。
        try:
            with tempfile.TemporaryDirectory() as tmp_dir:
                decrypted_path = Path(tmp_dir) / f"decrypted{file_path.suffix}"
                if self._decrypt_via_password_candidates(file_path, decrypted_path):
                    wb = openpyxl.load_workbook(
                        str(decrypted_path), read_only=True, data_only=True
                    )
                    return self._extract_xlsx_docs(wb, file_path)
        except Exception:
            pass
        return [Document(
            text=f"[XLSX解析失敗: パスワード保護または破損の可能性] {file_path.name}",
            source_path=file_path,
            location="stub",
        )]

    def _parse_pptx(self, file_path: Path) -> list[Document]:
        try:
            from pptx import Presentation
        except ImportError:
            return [Document(text=f"[PPTX未解析: python-pptx未インストール] {file_path.name}",
                             source_path=file_path, location="stub")]
        try:
            prs = Presentation(str(file_path))
            docs = []
            for i, slide in enumerate(prs.slides):
                texts = []
                for shape in slide.shapes:
                    if getattr(shape, "has_table", False):
                        for row in shape.table.rows:
                            # cell.textは複数段落/ソフト改行で\n・\vを含みうるため空白に正規化する
                            cells = [
                                " ".join(cell.text.split())
                                for cell in row.cells
                                if cell.text.strip()
                            ]
                            if cells:
                                texts.append(" | ".join(cells))
                        continue
                    if hasattr(shape, "text") and shape.text.strip():
                        texts.append(shape.text)
                if texts:
                    docs.append(Document(
                        text="\n".join(texts),
                        source_path=file_path,
                        location=f"slide_{i+1}",
                    ))
        except Exception:
            return [Document(
                text=f"[PPTX解析失敗: パスワード保護または破損の可能性] {file_path.name}",
                source_path=file_path,
                location="stub",
            )]
        return docs or [Document(text="", source_path=file_path, location="empty")]
