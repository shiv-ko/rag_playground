"""画像・グラフの読解生成。_call_vlm() はAnthropic Vision APIを呼び出す。"""
from __future__ import annotations

import base64
import json
import os
import re
import unicodedata
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

from anthropic import Anthropic

from src.generator.confidence_gate import ConfidenceGate
from src.models import Answer
from src.retriever.question_file_scope import find_named_files, find_question_stem_matches

_MEDIA_TYPES = {
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".gif": "image/gif",
    ".webp": "image/webp",
}

_VISUAL_EXTENSIONS = frozenset({*(_MEDIA_TYPES), ".pdf", ".pptx"})
_MAX_VISUAL_IMAGES = 80
_MAX_VISUAL_BYTES = 20 * 1024 * 1024
_MAX_SINGLE_VISUAL_BYTES = 5 * 1024 * 1024


@dataclass(frozen=True)
class VisualImage:
    data: bytes
    media_type: str
    label: str
    source_path: Path | None = None


def select_visual_images(images: list[VisualImage]) -> list[VisualImage]:
    """base64化後のAPIリクエスト上限に余裕を残して送信画像を確定する。"""
    selected: list[VisualImage] = []
    total_bytes = 0
    for visual in images:
        if len(selected) >= _MAX_VISUAL_IMAGES:
            break
        if len(visual.data) > _MAX_SINGLE_VISUAL_BYTES:
            continue
        if total_bytes + len(visual.data) > _MAX_VISUAL_BYTES:
            break
        selected.append(visual)
        total_bytes += len(visual.data)
    return selected

VLM_SYSTEM_PROMPT = """\
あなたは画像として提供されたグラフ・図表を読み取り、質問に答えるアシスタントです。

【重要ルール】
1. 画像に実際に描画されている内容のみを根拠に回答すること。読み取れない場合は正直に「わかりません」と答える
2. confidenceは「この回答が正しい確率」を0.0〜1.0で表すこと。読み取りに自信が持てない場合は0.4未満にすること
   （わからないのにconfidenceを高くすることは禁止）
3. 設問が特定の数値・日付・ラベルを問う場合、"answer" には値そのものだけを入れ、説明文を続けないこと
4. 回答は1000トークン以内に収めること

【出力形式】
{
  "answer": "回答テキスト",
  "confidence": 0.0〜1.0,
  "reasoning": "画像のどの部分から読み取ったか"
}
"""


class VLMImageAnswerer:
    """独立画像ファイル（png等）をClaude Vision APIに直接渡して質問に答える。"""

    def __init__(self, threshold: float = 0.4) -> None:
        self.gate = ConfidenceGate(threshold=threshold)
        self._client: Anthropic | None = None

    def answer(self, question: str, image_path: Path) -> Answer:
        media_type = _MEDIA_TYPES.get(image_path.suffix.lower())
        if media_type is None:
            return Answer(text=self.gate.missing_text(), confidence=0.0, was_gated=True,
                           gate_reason="vlm_unsupported_format")
        try:
            image_bytes = image_path.read_bytes()
        except OSError:
            return Answer(text=self.gate.missing_text(), confidence=0.0, was_gated=True,
                           gate_reason="vlm_file_not_found")

        image_b64 = base64.standard_b64encode(image_bytes).decode("ascii")
        raw = self._call_vlm(question, image_b64, media_type)
        answer_text, confidence, _reasoning = self._parse_response(raw)

        if not answer_text or not self.gate.should_answer(confidence):
            return Answer(text=self.gate.missing_text(), confidence=confidence, was_gated=True,
                           raw_text=answer_text, gate_reason="vlm_confidence")

        return Answer(text=answer_text, confidence=confidence, was_gated=False,
                      raw_text=answer_text, gate_reason="vlm_image")

    def answer_visuals(self, question: str, images: list[VisualImage]) -> Answer:
        """画像専用Office/PDFから抽出した複数画像を、出典順を保って一度に読解する。"""
        if not images:
            return Answer(text=self.gate.missing_text(), confidence=0.0, was_gated=True,
                          gate_reason="vlm_no_images")
        selected = select_visual_images(images)
        if not selected:
            return Answer(text=self.gate.missing_text(), confidence=0.0, was_gated=True,
                          gate_reason="vlm_image_limits")
        raw = self._call_vlm_images(question, selected)
        answer_text, confidence, _reasoning = self._parse_response(raw)
        if not answer_text or not self.gate.should_answer(confidence):
            return Answer(text=self.gate.missing_text(), confidence=confidence, was_gated=True,
                          raw_text=answer_text, gate_reason="vlm_confidence")
        return Answer(text=answer_text, confidence=confidence, was_gated=False,
                      raw_text=answer_text, gate_reason="vlm_document")

    def _get_client(self) -> Anthropic:
        if self._client is None:
            self._client = Anthropic(api_key=os.environ["ANTHROPIC_API_KEY"])
        return self._client

    def _call_vlm(self, question: str, image_b64: str, media_type: str) -> str:
        model = os.environ.get("CLAUDE_MODEL", "claude-sonnet-5")
        message = self._get_client().messages.create(
            model=model,
            max_tokens=1000,
            system=VLM_SYSTEM_PROMPT,
            messages=[{
                "role": "user",
                "content": [
                    {
                        "type": "image",
                        "source": {
                            "type": "base64",
                            "media_type": media_type,
                            "data": image_b64,
                        },
                    },
                    {"type": "text", "text": f"【質問】\n{question}"},
                ],
            }],
        )
        return "".join(block.text for block in message.content if hasattr(block, "text"))

    def _call_vlm_images(self, question: str, images: list[VisualImage]) -> str:
        content: list[dict] = []
        for visual in images:
            content.append({"type": "text", "text": f"【出典】{visual.label}"})
            content.append({
                "type": "image",
                "source": {
                    "type": "base64",
                    "media_type": visual.media_type,
                    "data": base64.standard_b64encode(visual.data).decode("ascii"),
                },
            })
        content.append({"type": "text", "text": f"【質問】\n{question}"})
        model = os.environ.get("CLAUDE_MODEL", "claude-sonnet-5")
        message = self._get_client().messages.create(
            model=model,
            max_tokens=1000,
            system=VLM_SYSTEM_PROMPT,
            messages=[{"role": "user", "content": content}],
        )
        return "".join(block.text for block in message.content if hasattr(block, "text"))

    def _parse_response(self, raw: str) -> tuple[str, float, str]:
        try:
            m = re.search(r"\{.*\}", raw, re.DOTALL)
            if m:
                data = json.loads(m.group(), strict=False)
                return (
                    str(data.get("answer", "")),
                    float(data.get("confidence", 0.0)),
                    str(data.get("reasoning", "")),
                )
        except (json.JSONDecodeError, ValueError, TypeError):
            pass
        return raw, 0.0, ""


# office_chart.pyの_FILE_REと同じ理由（\wが平仮名を含みUnicodeの助詞境界を吸収してしまう）で、
# 明示的な文字クラス（平仮名を除くASCII+漢字+カタカナ）を使う。
_IMAGE_FILE_RE = re.compile(r"([A-Za-z0-9_\-一-龠ァ-ヶー]+\.(?:png|jpg|jpeg|gif|webp))", re.IGNORECASE)


def find_referenced_image(question: str, project_name: str, data_dir: Path) -> Path | None:
    """質問文中に明示されたファイル名（例: figure_06.png）を案件ディレクトリ配下から探す。"""
    match = _IMAGE_FILE_RE.search(question)
    if match is None:
        return None
    filename = match.group(1)
    normalized_project = unicodedata.normalize("NFC", project_name)
    matching = sorted(
        p for p in data_dir.rglob(filename)
        if normalized_project in unicodedata.normalize("NFC", str(p))
    )
    return matching[0] if matching else None


def _media_type(data: bytes, name: str) -> str | None:
    suffix_type = _MEDIA_TYPES.get(Path(name).suffix.lower())
    if suffix_type is not None:
        return suffix_type
    try:
        from PIL import Image

        with Image.open(BytesIO(data)) as image:
            return Image.MIME.get(image.format or "")
    except Exception:
        return None


def _image_area(data: bytes) -> int:
    try:
        from PIL import Image

        with Image.open(BytesIO(data)) as image:
            return image.width * image.height
    except Exception:
        return 0


def extract_visual_images(path: Path) -> list[VisualImage]:
    """画像ファイル、画像専用PPTX、スキャンPDFからVLM投入用画像を抽出する。"""
    suffix = path.suffix.lower()
    if suffix in _MEDIA_TYPES:
        try:
            return [VisualImage(path.read_bytes(), _MEDIA_TYPES[suffix], path.name, path)]
        except OSError:
            return []
    if suffix == ".pptx":
        try:
            from pptx import Presentation

            presentation = Presentation(str(path))
            if any(
                getattr(shape, "has_text_frame", False) and shape.text.strip()
                for slide in presentation.slides for shape in slide.shapes
            ):
                return []
            result = []
            for slide_number, slide in enumerate(presentation.slides, 1):
                for shape in slide.shapes:
                    if not hasattr(shape, "image"):
                        continue
                    data = shape.image.blob
                    media_type = _media_type(data, shape.image.filename)
                    if media_type in _MEDIA_TYPES.values():
                        result.append(VisualImage(
                            data, media_type, f"{path.name} slide {slide_number}", path
                        ))
            return result
        except Exception:
            return []
    if suffix == ".pdf":
        try:
            import pypdf

            reader = pypdf.PdfReader(str(path))
            text_chars = sum(len((page.extract_text() or "").strip()) for page in reader.pages)
            if text_chars > max(100, len(reader.pages) * 10):
                return []
            result = []
            for page_number, page in enumerate(reader.pages, 1):
                candidates = []
                for image in page.images:
                    media_type = _media_type(image.data, image.name)
                    if media_type in _MEDIA_TYPES.values():
                        candidates.append((_image_area(image.data), image.data, media_type))
                if candidates:
                    _, data, media_type = max(candidates, key=lambda item: item[0])
                    result.append(VisualImage(
                        data, media_type, f"{path.name} page {page_number}", path
                    ))
            return result
        except Exception:
            return []
    return []


_SEMANTIC_FIGURES = {
    "feature_correlation_heatmap.png": ("相関", "correlation"),
    "target_distribution.png": ("ターゲット分布", "目的変数分布"),
    "missing_rate_top20.png": ("欠損率", "欠損割合"),
    "date_feature_trend.png": ("日付推移", "時系列推移"),
    "numeric_distribution_top6.png": ("数値分布", "数値特徴量分布"),
    "categorical_distribution_top3.png": ("カテゴリ分布", "カテゴリ特徴量分布"),
}


def find_referenced_visual_files(
    question: str, project_name: str | None, data_dir: Path
) -> list[Path]:
    """質問のファイル名・文書種別・共通figure規約から画像専用候補を絞り込む。"""
    normalized_question = unicodedata.normalize("NFC", question)
    normalized_project = unicodedata.normalize("NFC", project_name or "")
    candidates = sorted(
        path for path in data_dir.rglob("*")
        if path.is_file() and path.suffix.lower() in _VISUAL_EXTENSIONS
        and (
            (normalized_project and normalized_project in unicodedata.normalize("NFC", str(path)))
            or "社内管理" in unicodedata.normalize("NFC", str(path))
        )
    )
    names = [path.name for path in candidates]
    matched_names = find_named_files(normalized_question, names)
    if not matched_names:
        matched_names = find_question_stem_matches(normalized_question, names)
    if matched_names:
        matched_keys = {
            unicodedata.normalize("NFC", name).casefold() for name in matched_names
        }
        matched = [
            path for path in candidates
            if unicodedata.normalize("NFC", path.name).casefold() in matched_keys
        ]
    else:
        semantic_names = [
            name for name, hints in _SEMANTIC_FIGURES.items()
            if any(hint.casefold() in normalized_question.casefold() for hint in hints)
        ]
        if semantic_names:
            matched = [path for path in candidates if path.name in semantic_names]
        else:
            document_terms = []
            if "座席表" in normalized_question:
                document_terms.append("座席表")
            if (
                "会議録" in normalized_question
                or "アクション" in normalized_question
                or re.search(r"(?<![A-Za-z])AI(?![A-Za-z])", normalized_question)
            ):
                document_terms.append("会議録")
            if "報告資料" in normalized_question:
                document_terms.append("報告資料")
            if "最終報告" in normalized_question:
                document_terms.append("最終報告")
            matched = [path for path in candidates if any(term in path.stem for term in document_terms)]
    date_match = re.search(r"(\d{1,2})月(\d{1,2})日", normalized_question)
    if date_match and matched:
        month, day = (int(value) for value in date_match.groups())
        date_pattern = re.compile(rf"(?<!\d)0?{month}[-_/]0?{day}(?!\d)")
        dated = [
            path for path in matched
            if date_pattern.search(unicodedata.normalize("NFC", path.stem))
        ]
        if dated:
            matched = dated
    return [path for path in matched if extract_visual_images(path)]
