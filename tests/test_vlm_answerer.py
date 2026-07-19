from __future__ import annotations

import unicodedata
from pathlib import Path
from unittest.mock import patch

import pytest

from src.generator.vlm_answerer import (
    VLMImageAnswerer,
    VisualImage,
    extract_visual_images,
    find_referenced_visual_files,
    select_visual_images,
)

_PNG_1PX = bytes.fromhex(
    "89504e470d0a1a0a0000000d49484452000000010000000108020000009077"
    "53de0000000c4944415408d763f8ffff3f0005fe02fea739669f0000000049454e44ae426082"
)


class TestVLMImageAnswererCallConstruction:
    def test_sends_image_block_and_question(self, tmp_path: Path, monkeypatch) -> None:
        monkeypatch.setenv("ANTHROPIC_API_KEY", "test-key")
        monkeypatch.setenv("CLAUDE_MODEL", "claude-sonnet-5")
        image_path = tmp_path / "figure_06.png"
        image_path.write_bytes(_PNG_1PX)

        fake_content = type("C", (), {
            "text": '{"answer": "20\\u65e5", "confidence": 0.85, "reasoning": "r"}'
        })()
        fake_response = type("R", (), {"content": [fake_content]})()

        answerer = VLMImageAnswerer()
        with patch("src.generator.vlm_answerer.Anthropic") as MockAnthropic:
            MockAnthropic.return_value.messages.create.return_value = fake_response
            result = answerer.answer("dayによる件数推移で件数が最も高いのは何日ですか。", image_path)

        assert result.was_gated is False
        assert result.text == "20日"

        _, kwargs = MockAnthropic.return_value.messages.create.call_args
        assert kwargs["model"] == "claude-sonnet-5"
        content_blocks = kwargs["messages"][0]["content"]
        image_blocks = [b for b in content_blocks if b["type"] == "image"]
        text_blocks = [b for b in content_blocks if b["type"] == "text"]
        assert len(image_blocks) == 1
        assert image_blocks[0]["source"]["media_type"] == "image/png"
        assert "件数が最も高い" in text_blocks[0]["text"]

    def test_low_confidence_is_gated(self, tmp_path: Path, monkeypatch) -> None:
        monkeypatch.setenv("ANTHROPIC_API_KEY", "test-key")
        image_path = tmp_path / "figure_06.png"
        image_path.write_bytes(_PNG_1PX)

        fake_content = type("C", (), {
            "text": '{"answer": "わかりません", "confidence": 0.1, "reasoning": "r"}'
        })()
        fake_response = type("R", (), {"content": [fake_content]})()

        answerer = VLMImageAnswerer()
        with patch("src.generator.vlm_answerer.Anthropic") as MockAnthropic:
            MockAnthropic.return_value.messages.create.return_value = fake_response
            result = answerer.answer("質問", image_path)

        assert result.was_gated is True

    def test_missing_file_is_gated_without_api_call(self, tmp_path: Path, monkeypatch) -> None:
        monkeypatch.setenv("ANTHROPIC_API_KEY", "test-key")
        answerer = VLMImageAnswerer()

        with patch("src.generator.vlm_answerer.Anthropic") as MockAnthropic:
            result = answerer.answer("質問", tmp_path / "not_exist.png")

        assert result.was_gated is True
        MockAnthropic.return_value.messages.create.assert_not_called()

    def test_unsupported_extension_is_gated_without_api_call(self, tmp_path: Path, monkeypatch) -> None:
        monkeypatch.setenv("ANTHROPIC_API_KEY", "test-key")
        file_path = tmp_path / "report.pdf"
        file_path.write_bytes(b"%PDF-1.4")
        answerer = VLMImageAnswerer()

        with patch("src.generator.vlm_answerer.Anthropic") as MockAnthropic:
            result = answerer.answer("質問", file_path)

        assert result.was_gated is True
        MockAnthropic.return_value.messages.create.assert_not_called()

    def test_non_json_response_is_gated_gracefully(self, tmp_path: Path, monkeypatch) -> None:
        monkeypatch.setenv("ANTHROPIC_API_KEY", "test-key")
        image_path = tmp_path / "figure_06.png"
        image_path.write_bytes(_PNG_1PX)

        fake_content = type("C", (), {"text": "申し訳ありませんが、画像を読み取れませんでした。"})()
        fake_response = type("R", (), {"content": [fake_content]})()

        answerer = VLMImageAnswerer()
        with patch("src.generator.vlm_answerer.Anthropic") as MockAnthropic:
            MockAnthropic.return_value.messages.create.return_value = fake_response
            result = answerer.answer("質問", image_path)

        assert result.was_gated is True

    def test_null_confidence_in_response_is_gated_not_crashed(self, tmp_path: Path, monkeypatch) -> None:
        monkeypatch.setenv("ANTHROPIC_API_KEY", "test-key")
        image_path = tmp_path / "figure_06.png"
        image_path.write_bytes(_PNG_1PX)

        fake_content = type("C", (), {
            "text": '{"answer": "20\\u65e5", "confidence": null, "reasoning": "r"}'
        })()
        fake_response = type("R", (), {"content": [fake_content]})()

        answerer = VLMImageAnswerer()
        with patch("src.generator.vlm_answerer.Anthropic") as MockAnthropic:
            MockAnthropic.return_value.messages.create.return_value = fake_response
            result = answerer.answer("質問", image_path)

        assert result.was_gated is True


def test_extracts_picture_blob_from_image_only_pptx(tmp_path: Path) -> None:
    from io import BytesIO

    from pptx import Presentation
    from pptx.util import Inches

    image_path = tmp_path / "seat.png"
    image_path.write_bytes(_PNG_1PX)
    pptx_path = tmp_path / "座席表.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(str(image_path), Inches(1), Inches(1))
    prs.save(pptx_path)

    images = extract_visual_images(pptx_path)

    assert len(images) == 1
    assert images[0].media_type == "image/png"
    assert images[0].label == "座席表.pptx slide 1"
    assert BytesIO(images[0].data).read(8) == b"\x89PNG\r\n\x1a\n"


def test_extracts_largest_image_from_each_scanned_pdf_page(tmp_path: Path) -> None:
    from PIL import Image

    pdf_path = tmp_path / "会議録.pdf"
    pages = [Image.new("RGB", (80, 60), color) for color in ("white", "black")]
    pages[0].save(pdf_path, save_all=True, append_images=pages[1:])

    images = extract_visual_images(pdf_path)

    assert [image.label for image in images] == [
        "会議録.pdf page 1",
        "会議録.pdf page 2",
    ]
    assert all(image.media_type in {"image/jpeg", "image/png"} for image in images)


def test_find_referenced_visual_files_uses_stem_and_image_only_gate(tmp_path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    project = "テスト株式会社"
    project_dir = tmp_path / "プロジェクト" / project
    project_dir.mkdir(parents=True)
    image_path = tmp_path / "seat.png"
    image_path.write_bytes(_PNG_1PX)
    pptx_path = project_dir / "座席表.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(str(image_path), Inches(1), Inches(1))
    prs.save(pptx_path)

    assert find_referenced_visual_files("テストの座席表で右隣は誰ですか", project, tmp_path) == [
        pptx_path
    ]


def test_find_referenced_visual_files_uses_common_semantic_figure(tmp_path: Path) -> None:
    project = "テスト株式会社"
    figure = tmp_path / "プロジェクト" / project / "reports" / "figures" / "feature_correlation_heatmap.png"
    figure.parent.mkdir(parents=True)
    figure.write_bytes(_PNG_1PX)

    result = find_referenced_visual_files(
        "テストの特徴量のうちターゲットとの相関が最大なのは何ですか", project, tmp_path
    )

    assert result == [figure]


def test_find_referenced_visual_files_normalizes_nfd_name_and_filters_date(tmp_path: Path) -> None:
    from PIL import Image

    project = "テスト株式会社"
    project_dir = tmp_path / "プロジェクト" / project
    project_dir.mkdir(parents=True)
    wanted = project_dir / unicodedata.normalize("NFD", "報告資料_2025-05-27.pdf")
    other = project_dir / "報告資料_2025-06-17.pdf"
    Image.new("RGB", (20, 20), "white").save(wanted)
    Image.new("RGB", (20, 20), "white").save(other)

    result = find_referenced_visual_files(
        "テストの5月27日の報告資料で確認してください", project, tmp_path
    )

    assert result == [wanted]


def test_find_referenced_visual_files_routes_action_item_to_meeting_minutes(tmp_path: Path) -> None:
    from PIL import Image

    project = "テスト株式会社"
    meeting = tmp_path / "プロジェクト" / project / "会議録_2025-04-03.pdf"
    meeting.parent.mkdir(parents=True)
    Image.new("RGB", (20, 20), "white").save(meeting)

    result = find_referenced_visual_files(
        "テストのM01からM02までに完了したAIを答えてください", project, tmp_path
    )

    assert result == [meeting]


def test_text_bearing_pdf_is_not_treated_as_scanned_visual(tmp_path: Path) -> None:
    import pypdf

    pdf_path = tmp_path / "通常報告.pdf"
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.add_metadata({"/Subject": "x" * 300})
    with pdf_path.open("wb") as stream:
        writer.write(stream)

    # 空白ページ自体は画像を持たないため、安全側に候補外となる。
    assert extract_visual_images(pdf_path) == []


def test_answer_visuals_sends_all_images_with_source_labels() -> None:
    answerer = VLMImageAnswerer()
    captured = {}
    answerer._call_vlm_images = lambda question, images: captured.update(  # type: ignore[method-assign]
        question=question, images=images
    ) or '{"answer":"A10の内容","confidence":0.9,"reasoning":"page 2"}'

    result = answerer.answer_visuals("A10の内容は？", [
        VisualImage(_PNG_1PX, "image/png", "会議録.pdf page 1"),
        VisualImage(_PNG_1PX, "image/png", "会議録.pdf page 2"),
    ])

    assert result.was_gated is False
    assert result.text == "A10の内容"
    assert [image.label for image in captured["images"]] == [
        "会議録.pdf page 1", "会議録.pdf page 2"
    ]


def test_select_visual_images_skips_oversized_first_image() -> None:
    oversized = VisualImage(b"x" * (5 * 1024 * 1024 + 1), "image/png", "oversized")
    small = VisualImage(_PNG_1PX, "image/png", "small")

    assert select_visual_images([oversized, small]) == [small]
