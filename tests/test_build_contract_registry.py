"""contracts.jsonl extraction tests."""
from __future__ import annotations

import io
import json
from pathlib import Path

import msoffcrypto
from docx import Document as DocxDocument
from pptx import Presentation
from pptx.util import Inches

from scripts.build_contract_registry import (
    build_contract_registry,
    candidate_dates_from_filename,
    classify_rounding_rule,
    extract_dates,
    extract_payment_schedule,
    extract_report_values,
    literal_password_from_filename,
    parse_contract_text,
)
from src.utils.office_crypto import derive_office_password


def test_parse_time_and_materials_contract_values() -> None:
    text = """
5. 契約期間
本契約の契約期間は、2025-07-08から2025-08-11までの5週間とする。
6. 報酬および支払条件
本契約の料金モデルは、time_and_materialsとし、実績工数に基づく事後精算（月次精算）とする。
時間単価は25,000円（消費税別）とする。
想定総工数は170時間とする。
見込金額は、税抜4,250,000円、消費税425,000円、税込4,675,000円とする。
作業時間の計上単位は30分とし、30分未満の端数は30分単位に切り上げて計上する。
7. 知的財産
"""
    row = parse_contract_text("テスト社", "契約書.docx", text)
    assert row["contract_type"] == "time_and_materials"
    assert row["rate_yen_per_hour"] == 25000
    assert row["esth_hours"] == 170
    assert row["estimated_amount_incl_tax"] == 4675000
    assert row["rounding_rule"] == "ceiling_to_unit"
    assert row["contract_period_days"] == 35


def test_parse_time_and_materials_contract_values_tax_status_first_bullet_form() -> None:
    # 実データ(かえで総合病院)で確認された表記ゆれ: 「見込金額（税込）」のような
    # 名目→税区分の順ではなく、「税込見込金額」のように税区分→名目の順で、
    # かつ箇条書き(- ラベル：値)形式になっているケース。
    text = """
5. 契約期間
本契約の契約期間は、2025-09-02から2025-10-07までの5週間とする。
6.1 料金体系
本契約の料金体系は、time_and_materialsとし、以下の条件を適用する。
- 通貨：JPY
- 請求単位：hour
- 時間単価：25,000円
- 想定総工数：140時間
- 税抜見込金額：3,500,000円
- 消費税率：10%
- 消費税額：350,000円
- 税込見込金額：3,850,000円
"""
    row = parse_contract_text("テスト社2", "契約書.docx", text)
    assert row["contract_type"] == "time_and_materials"
    assert row["estimated_amount_excl_tax"] == 3500000
    assert row["estimated_amount_incl_tax"] == 3850000


def test_parse_fixed_contract_values() -> None:
    text = """
5. 契約期間
本契約の契約期間は、2025-10-01から2025-11-11までの6週間とする。
6. 報酬および支払条件
本契約の契約形態は固定価格契約とし、契約金額（税抜）：5,250,000円、
消費税額：525,000円、契約金額（税込）：5,775,000円とする。
支払計画: 着手金50%（期日: 2025-10-08）＋ 検収金50%。
"""
    row = parse_contract_text("固定社", "契約書.docx", text)
    assert row["contract_type"] == "fixed"
    assert row["estimated_amount_incl_tax"] == 5775000
    assert row["advance_payment_amount"] == 2887500


def test_classify_half_unit_bucket_rounding() -> None:
    clause = "工数計上の丸め単位は30分とし、30分未満を0.5時間、30分超60分未満を1.0時間として0.5時間単位で計上する。"
    assert classify_rounding_rule(clause) == "half_unit_bucket"


def test_parse_advance_payment_from_docx_table_row_with_percent_column() -> None:
    # 実データの表形式（京橋信用ソリューションズと同じ列順）:
    # 支払回 | 名目 | 比率 | 金額（税抜） | 消費税額 | 金額（税込） | 支払条件 | 支払期日
    text = """
6. 報酬および支払条件
6.1 契約金額
本契約の契約形態は固定価格契約とし、契約金額（税抜）：5,250,000円、
消費税額：525,000円、契約金額（税込）：5,775,000円とする。
6.2 支払条件
支払回 | 名目 | 比率 | 金額（税抜） | 消費税額 | 金額（税込） | 支払条件 | 支払期日
第1回 | 着手金 | 50% | 2,625,000円 | 262,500円 | 2,887,500円 | 契約締結後5営業日以内 | 2025-10-08
第2回 | 検収金 | 50% | 2,625,000円 | 262,500円 | 2,887,500円 | 検収完了後5営業日以内 | 2025-11-19
"""
    row = parse_contract_text("京橋風", "契約書.docx", text)
    assert row["advance_payment_amount"] == 2_887_500


def test_parse_advance_payment_from_docx_table_row_with_different_column_order() -> None:
    # 実データの表形式（青葉与信マネジメントと同じ列順、支払条件が税抜より前に来る）:
    # 支払回 | 支払名目 | 比率 | 支払条件 | 税抜金額 | 消費税額 | 税込金額 | 支払期日
    text = """
6. 報酬および支払条件
本契約の契約形態は固定価格とし、契約時に金額を固定し、工数実績による事後精算は行わない。
本契約の報酬総額は、税抜4,200,000円、消費税420,000円、税込4,620,000円とする。
支払回 | 支払名目 | 比率 | 支払条件 | 税抜金額 | 消費税額 | 税込金額 | 支払期日
1 | 着手金 | 50% | 契約締結後5営業日以内 | 2,100,000円 | 210,000円 | 2,310,000円 | 2025-04-16
2 | 検収金 | 50% | 検収完了後5営業日以内 | 2,100,000円 | 210,000円 | 2,310,000円 | 2025-06-04
"""
    row = parse_contract_text("青葉与信風", "契約書.docx", text)
    assert row["advance_payment_amount"] == 2_310_000


def test_extract_dates_allows_whitespace_before_kara_and_made() -> None:
    # 実データで見つかった表記ゆれ:「日 から」「日 まで」のように日付とキーワードの間にスペースが入る。
    text = "本契約の契約期間は、2025-05-13 から 2025-07-22 まで とする。"
    start, end, days = extract_dates(text)
    assert start == "2025-05-13"
    assert end == "2025-07-22"
    assert days == 71


def test_extract_dates_no_space_form_still_parses() -> None:
    # 既存の非スペース表記（従来通り）が引き続き通ることの非回帰確認。
    text = "本契約の契約期間は、2025-07-08から2025-08-11までの5週間とする。"
    start, end, days = extract_dates(text)
    assert start == "2025-07-08"
    assert end == "2025-08-11"
    assert days == 35


# --- extract_report_values / has_final_report (06.報告書 file-existence flag) ---


def _make_pptx_with_text(path: Path, text: str) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(2))
    textbox.text_frame.text = text
    prs.save(path)


def test_extract_report_values_true_when_non_old_report_file_exists(tmp_path: Path) -> None:
    # 実データ(青葉与信・固定価格)を模した回帰ケース: 最終報告書の金額表記は
    # 「最終請求金額（税込）」「税込金額」のどちらの正規表現にも一致しないが、
    # 報告ファイル自体は存在する＝完了とみなすべき。
    project_dir = tmp_path / "テスト案件"
    report_dir = project_dir / "06.報告書"
    report_dir.mkdir(parents=True)
    _make_pptx_with_text(
        report_dir / "最終報告.pptx", "契約金額：¥4,200,000（税抜）/ ¥4,620,000（税込）"
    )
    values = extract_report_values(project_dir)
    assert values["has_final_report"] is True


def test_extract_report_values_false_when_only_old_report_file_exists(tmp_path: Path) -> None:
    project_dir = tmp_path / "テスト案件2"
    report_dir = project_dir / "06.報告書"
    report_dir.mkdir(parents=True)
    _make_pptx_with_text(report_dir / "最終報告_old.pptx", "旧版のダミーテキスト")
    values = extract_report_values(project_dir)
    assert values["has_final_report"] is False


def test_extract_report_values_false_when_no_report_dir(tmp_path: Path) -> None:
    project_dir = tmp_path / "テスト案件3"
    project_dir.mkdir(parents=True)
    values = extract_report_values(project_dir)
    assert values["has_final_report"] is False


# --- office_crypto wiring (encrypted contract fallback) ---

# msoffcrypto-tool 6.0.0のOLEコンテナ書き込み処理は、暗号化payloadが小さい(<=4096バイト)と
# mini-FAT/regular-FATの不整合で往復（自前暗号化→自前復号）してもバイト列が壊れることがある
# （tests/test_office_crypto.pyのコメントで裏取り済みの既知のクセ）。ここでは実際に
# python-docxで開ける最小限のdocxを作り、十分な段落を足して4KB超のペイロードにしてから
# 暗号化することでこの不具合を回避する。
def _build_real_docx_bytes(paragraphs: list[str]) -> bytes:
    doc = DocxDocument()
    for text in paragraphs:
        doc.add_paragraph(text)
    for i in range(300):
        doc.add_paragraph(f"padding padding padding padding padding paragraph {i}")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _write_encrypted_docx(dest: Path, password: str, paragraphs: list[str]) -> None:
    plain_buf = io.BytesIO(_build_real_docx_bytes(paragraphs))
    office_file = msoffcrypto.OfficeFile(plain_buf)
    dest.parent.mkdir(parents=True, exist_ok=True)
    with open(dest, "wb") as f:
        office_file.encrypt(password, f)


def _write_project_registry(path: Path, project_name: str, primary_alias: str) -> None:
    path.write_text(
        json.dumps([{"project_name": project_name, "primary_alias": primary_alias}], ensure_ascii=False),
        encoding="utf-8",
    )


def test_candidate_dates_from_filename_extracts_8digit_date_after_pw_prefix() -> None:
    # 実運用の命名慣習（docs/encrypted_file_queue.md記載）: `pw-<略号><開始年月日8桁>`。
    path = Path("契約書_pw-testalias20250115.docx")
    assert candidate_dates_from_filename(path) == ["20250115"]


def test_candidate_dates_from_filename_empty_when_no_pw_marker() -> None:
    assert candidate_dates_from_filename(Path("契約書.docx")) == []


def test_literal_password_from_filename_extracts_token_after_pw_prefix() -> None:
    # 実データ(かえで案件)で判明した実運用パターン: `pw-<トークン>`のトークン自体が
    # DA-規則を介さずそのまま平文パスワードになっているケースがある。
    path = Path("契約書_pw-kaede20250902.docx")
    assert literal_password_from_filename(path) == "kaede20250902"


def test_literal_password_from_filename_none_when_no_pw_marker() -> None:
    assert literal_password_from_filename(Path("契約書.docx")) is None


def test_build_contract_registry_decrypts_via_filename_date_candidate(tmp_path: Path) -> None:
    project_root = tmp_path / "projects"
    contract_dir = project_root / "検証用医療法人テスト" / "01.契約"
    contract_dir.mkdir(parents=True)

    alias = "TESTALIAS"
    password = derive_office_password(alias, "2025-01-15", ".docx")
    _write_encrypted_docx(
        contract_dir / "契約書_pw-testalias20250115.docx",
        password,
        [
            "5. 契約期間",
            "本契約の契約期間は、2025-01-15から2025-02-11までの4週間とする。",
            "6. 報酬および支払条件",
            "本契約の契約形態は固定価格契約とし、契約金額（税抜）：1,000,000円、"
            "消費税額：100,000円、契約金額（税込）：1,100,000円とする。",
        ],
    )

    project_registry_path = tmp_path / "project_registry.json"
    _write_project_registry(project_registry_path, "検証用医療法人テスト", alias)

    rows = build_contract_registry(
        project_root=project_root,
        project_registry_path=project_registry_path,
        schedule_tasks_path=tmp_path / "no_schedule.jsonl",
    )

    assert len(rows) == 1
    row = rows[0]
    assert row["status"] == "ok"
    assert row["start_date"] == "2025-01-15"
    assert row["end_date"] == "2025-02-11"
    assert row["estimated_amount_incl_tax"] == 1100000


def test_build_contract_registry_decrypts_via_literal_filename_password(tmp_path: Path) -> None:
    # 実データ(かえで案件)で判明したパターン: `pw-<トークン>`のトークンがDA-規則を介さず
    # そのまま平文パスワードになっている（DA-規則の候補は全滅する状況でも、こちらで復号できる）。
    project_root = tmp_path / "projects"
    contract_dir = project_root / "検証用医療法人テスト6" / "01.契約"
    contract_dir.mkdir(parents=True)

    literal_password = "kaede20250902"
    _write_encrypted_docx(
        contract_dir / "契約書_pw-kaede20250902.docx",
        literal_password,
        [
            "5. 契約期間",
            "本契約の契約期間は、2025-09-02から2025-10-14までの6週間とする。",
            "6. 報酬および支払条件",
            "本契約の契約形態は固定価格契約とし、契約金額（税抜）：3,000,000円、"
            "消費税額：300,000円、契約金額（税込）：3,300,000円とする。",
        ],
    )

    project_registry_path = tmp_path / "project_registry.json"
    _write_project_registry(project_registry_path, "検証用医療法人テスト6", "KAEDE")

    rows = build_contract_registry(
        project_root=project_root,
        project_registry_path=project_registry_path,
        schedule_tasks_path=tmp_path / "no_schedule.jsonl",
    )

    assert len(rows) == 1
    row = rows[0]
    assert row["status"] == "ok"
    assert row["start_date"] == "2025-09-02"
    assert row["estimated_amount_incl_tax"] == 3300000


def test_build_contract_registry_falls_back_to_schedule_date_candidate_when_filename_has_none(
    tmp_path: Path,
) -> None:
    project_root = tmp_path / "projects"
    contract_dir = project_root / "検証用医療法人テスト2" / "01.契約"
    contract_dir.mkdir(parents=True)

    alias = "TESTALIAS2"
    password = derive_office_password(alias, "2025-03-03", ".docx")
    # ファイル名には日付候補が無い(pw-マーカー無し)ので、スケジュールregistry側の候補で試す。
    _write_encrypted_docx(
        contract_dir / "契約書.docx",
        password,
        [
            "5. 契約期間",
            "本契約の契約期間は、2025-03-03から2025-03-31までの4週間とする。",
            "6. 報酬および支払条件",
            "本契約の契約形態は固定価格契約とし、契約金額（税抜）：2,000,000円、"
            "消費税額：200,000円、契約金額（税込）：2,200,000円とする。",
        ],
    )

    project_registry_path = tmp_path / "project_registry.json"
    _write_project_registry(project_registry_path, "検証用医療法人テスト2", alias)

    schedule_tasks_path = tmp_path / "schedule_tasks.jsonl"
    schedule_row = {
        "project_name": "検証用医療法人テスト2",
        "values": {"開始日": "2025-03-03T00:00:00", "終了日": "2025-03-31T00:00:00"},
    }
    schedule_tasks_path.write_text(json.dumps(schedule_row, ensure_ascii=False) + "\n", encoding="utf-8")

    rows = build_contract_registry(
        project_root=project_root,
        project_registry_path=project_registry_path,
        schedule_tasks_path=schedule_tasks_path,
    )

    assert len(rows) == 1
    assert rows[0]["status"] == "ok"
    assert rows[0]["start_date"] == "2025-03-03"


def test_build_contract_registry_falls_back_to_failed_when_no_date_candidate_available(
    tmp_path: Path,
) -> None:
    project_root = tmp_path / "projects"
    contract_dir = project_root / "検証用医療法人テスト3" / "01.契約"
    contract_dir.mkdir(parents=True)

    password = derive_office_password("UNKNOWNALIAS", "2025-03-03", ".docx")
    _write_encrypted_docx(contract_dir / "契約書.docx", password, ["本文"])

    project_registry_path = tmp_path / "project_registry.json"
    _write_project_registry(project_registry_path, "検証用医療法人テスト3", "UNKNOWNALIAS")

    rows = build_contract_registry(
        project_root=project_root,
        project_registry_path=project_registry_path,
        schedule_tasks_path=tmp_path / "no_schedule.jsonl",
    )

    assert len(rows) == 1
    assert rows[0]["status"] == "failed"


def test_build_contract_registry_falls_back_to_failed_when_all_candidates_wrong(
    tmp_path: Path,
) -> None:
    project_root = tmp_path / "projects"
    contract_dir = project_root / "検証用医療法人テスト4" / "01.契約"
    contract_dir.mkdir(parents=True)

    real_password = derive_office_password("WRONGALIAS", "2025-06-06", ".docx")
    # ファイル名の日付候補(20250101)は実際のパスワード生成に使った日付(2025-06-06)と異なるため、
    # 候補はあるが全滅してInvalidKeyErrorとなり、既存どおりstatus=failedへ後退する。
    _write_encrypted_docx(contract_dir / "契約書_pw-wrongalias20250101.docx", real_password, ["本文"])

    project_registry_path = tmp_path / "project_registry.json"
    _write_project_registry(project_registry_path, "検証用医療法人テスト4", "WRONGALIAS")

    rows = build_contract_registry(
        project_root=project_root,
        project_registry_path=project_registry_path,
        schedule_tasks_path=tmp_path / "no_schedule.jsonl",
    )

    assert len(rows) == 1
    assert rows[0]["status"] == "failed"


def test_build_contract_registry_non_encrypted_corruption_still_falls_back_to_failed(
    tmp_path: Path,
) -> None:
    # 暗号化(CDFV2)ではない、単なる破損/非docxファイルは、復号フォールバックを試みずに
    # 従来どおりstatus=failedへ後退すること（今回の変更による非回帰の確認）。
    project_root = tmp_path / "projects"
    contract_dir = project_root / "検証用医療法人テスト5" / "01.契約"
    contract_dir.mkdir(parents=True)
    (contract_dir / "契約書.docx").write_bytes(b"not a docx at all, just garbage bytes")

    project_registry_path = tmp_path / "project_registry.json"
    _write_project_registry(project_registry_path, "検証用医療法人テスト5", "IRRELEVANT")

    rows = build_contract_registry(
        project_root=project_root,
        project_registry_path=project_registry_path,
        schedule_tasks_path=tmp_path / "no_schedule.jsonl",
    )

    assert len(rows) == 1
    assert rows[0]["status"] == "failed"


# --- extract_payment_schedule (Q40用: 支払回テーブルの月次集計) ---


def test_extract_payment_schedule_basic_two_rows() -> None:
    # 実データ（京橋信用ソリューションズ）の列順: 支払回|名目|比率|税抜|消費税額|税込|支払条件|支払期日
    text = """
支払回 | 名目 | 比率 | 金額（税抜） | 消費税額 | 金額（税込） | 支払条件 | 支払期日
第1回 | 着手金 | 50% | 2,625,000円 | 262,500円 | 2,887,500円 | 契約締結後5営業日以内 | 2025-10-08
第2回 | 検収金 | 50% | 2,625,000円 | 262,500円 | 2,887,500円 | 検収完了後5営業日以内 | 2025-11-19
"""
    schedule = extract_payment_schedule(text)
    assert schedule == [
        {"month": "2025-10", "amount_incl_tax": 2_887_500, "due_date": "2025-10-08"},
        {"month": "2025-11", "amount_incl_tax": 2_887_500, "due_date": "2025-11-19"},
    ]


def test_extract_payment_schedule_different_column_order() -> None:
    # 実データ（青葉与信マネジメント）の列順: 支払回|支払名目|比率|支払条件|税抜|消費税額|税込|支払期日
    text = """
支払回 | 支払名目 | 比率 | 支払条件 | 税抜金額 | 消費税額 | 税込金額 | 支払期日
1 | 着手金 | 50% | 契約締結後5営業日以内 | 2,100,000円 | 210,000円 | 2,310,000円 | 2025-04-16
2 | 検収金 | 50% | 検収完了後5営業日以内 | 2,100,000円 | 210,000円 | 2,310,000円 | 2025-06-04
"""
    schedule = extract_payment_schedule(text)
    assert schedule == [
        {"month": "2025-04", "amount_incl_tax": 2_310_000, "due_date": "2025-04-16"},
        {"month": "2025-06", "amount_incl_tax": 2_310_000, "due_date": "2025-06-04"},
    ]


def test_extract_payment_schedule_date_embedded_in_condition_cell() -> None:
    # 実データ（青嶺不動産アセットマネジメント）: 独立した支払期日列が無く、最終列（支払期限）に
    # 「条件文（YYYY-MM-DD）」の複合値として日付が埋め込まれている。
    text = """
支払回 | マイルストーン | 比率 | 金額（税抜） | 消費税額 | 金額（税込） | 支払期限
1 | 最終一括精算 | 100% | 4,250,000円 | 425,000円 | 4,675,000円 | 最終成果物の検収完了後5営業日以内（2025-09-24）
"""
    schedule = extract_payment_schedule(text)
    assert schedule == [
        {"month": "2025-09", "amount_incl_tax": 4_675_000, "due_date": "2025-09-24"},
    ]


def test_extract_payment_schedule_strips_mikomi_suffix_in_amount_cell() -> None:
    # 実データ（東都人材プラットフォーム、T&M最終一括精算）: 税込金額セルに「（見込）」が付く。
    text = """
支払回 | マイルストーン | 比率 | 金額（税抜） | 消費税額 | 金額（税込） | 支払条件 | 支払期日
1 | 最終一括精算 | 100% | 4,250,000円（見込） | 425,000円（見込） | 4,675,000円（見込） | 最終成果物の検収完了後5営業日以内 | 2025-10-06
"""
    schedule = extract_payment_schedule(text)
    assert schedule == [
        {"month": "2025-10", "amount_incl_tax": 4_675_000, "due_date": "2025-10-06"},
    ]


def test_extract_payment_schedule_empty_when_no_table() -> None:
    assert extract_payment_schedule("支払条件については別途協議する。") == []


def test_parse_contract_text_includes_payment_schedule() -> None:
    text = """
6. 報酬および支払条件
支払回 | 名目 | 比率 | 金額（税抜） | 消費税額 | 金額（税込） | 支払条件 | 支払期日
第1回 | 着手金 | 50% | 2,625,000円 | 262,500円 | 2,887,500円 | 契約締結後5営業日以内 | 2025-10-08
"""
    row = parse_contract_text("京橋風", "契約書.docx", text)
    assert row["payment_schedule"] == [
        {"month": "2025-10", "amount_incl_tax": 2_887_500, "due_date": "2025-10-08"},
    ]
