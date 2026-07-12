from __future__ import annotations

import json
from pathlib import Path

from src.generator.analysis_answerer import AnalysisAnswerer
from src.structured.artifact_store import StructuredArtifactStore
from src.utils.question_classifier import classify_question


def _store(tmp_path: Path, rows: list[dict]) -> StructuredArtifactStore:
    (tmp_path / "analysis_records.jsonl").write_text(
        "\n".join(json.dumps(row, ensure_ascii=False) for row in rows), encoding="utf-8"
    )
    return StructuredArtifactStore.from_artifacts_dir(tmp_path)


def _row(kind: str, payload: dict, path: str = "src/modeling.py") -> dict:
    return {"project_name": "A社", "kind": kind, "source_path": path, "payload": payload}


def test_sparse_output_false_model_is_derived_from_comparison(tmp_path: Path) -> None:
    store = _store(tmp_path, [_row("python", {"assignments": [], "comparisons": [], "calls": [
        {"function": "build_preprocessor", "args": [], "keywords": {
            "sparse_output": {"expression": "model_key != 'hist_gradient_boosting'"}
        }}
    ]})])
    answer = AnalysisAnswerer().answer(
        "modeling.pyで前処理器のsparse_outputがFalseになるmodel_typeは何ですか", "A社", store
    )
    assert answer is not None
    assert answer.text == "hist_gradient_boosting"
    assert answer.confidence == 0.95
    assert answer.source_docs[0].document.source_path == Path("src/modeling.py")


def test_cat_condition_reports_actual_dtype_candidates_and_operator(tmp_path: Path) -> None:
    payload = {
        "assignments": [],
        "comparisons": [{"left": "unique_count", "operators": ["GtE"], "comparators": [{"name": "categorical_unique_limit"}]}],
        "calls": [
            {"function": "pd.api.types.is_object_dtype", "args": [{"name": "series"}], "keywords": {}},
            {"function": "pd.api.types.is_string_dtype", "args": [{"name": "series"}], "keywords": {}},
            {"function": "pd.api.types.is_categorical_dtype", "args": [{"name": "series"}], "keywords": {}},
        ],
    }
    config = _row("json", {"feature_plan": {"categorical_unique_limit": 50}}, "configs/project_config.json")
    store = _store(tmp_path, [_row("python", payload, "src/features.py"), config])
    answer = AnalysisAnswerer().answer("CATはdtypeとユニーク数の条件でどのように判定しますか", "A社", store)
    assert answer is not None
    assert answer.text == "object・string・categorical型で、かつ unique_count < 50 の列をCATと判定します。"


def test_runtime_model_defaults_merge_code_and_config(tmp_path: Path) -> None:
    config = _row("json", {"model_type": "gradient_boosting", "model_params": {}, "random_state": 42}, "configs/project_config.json")
    code = _row("python", {
        "assignments": [
            {"target": "n_estimators", "value": {"expression": "to_int(model_params.get('n_estimators'), 300)"}},
            {"target": "learning_rate", "value": {"expression": "to_float(model_params.get('learning_rate'), 0.1)"}},
        ],
        "comparisons": [],
        "calls": [{"function": "GradientBoostingClassifier", "args": [], "keywords": {
            "n_estimators": {"name": "n_estimators"}, "learning_rate": {"name": "learning_rate"},
            "random_state": {"name": "random_state"}
        }}],
    })
    store = _store(tmp_path, [config, code])
    answer = AnalysisAnswerer().answer(
        "勾配ブースティング法に実際に渡されるn_estimators、learning_rate、random_stateはそれぞれいくつですか", "A社", store
    )
    assert answer is not None
    assert answer.text == "n_estimators=300、learning_rate=0.1、random_state=42"


def test_metrics_json_path_and_generated_interaction_columns(tmp_path: Path) -> None:
    metrics = _row("json", {
        "model_params": {"max_depth": 12},
        "feature_selection": {"selected_columns": ["a", "a__x__b", "b__x__c"]},
    }, "analysis_outputs/metrics.json")
    code = _row("python", {
        "assignments": [{"target": "feature_name", "value": {"expression": "f'{left}__x__{right}'"}}],
        "comparisons": [], "calls": [],
    }, "src/features.py")
    store = _store(tmp_path, [metrics, code])
    answer = AnalysisAnswerer().answer("最良モデルのパラメータmax_depthはいくらですか", "A社", store)
    assert answer is not None and answer.text == "12"
    listed = AnalysisAnswerer().answer(
        "feature_selection.selected_columnsに含まれる生成された数値交互作用特徴量をすべて答えてください", "A社", store
    )
    assert listed is not None and listed.text == "a__x__b、b__x__c"


def test_ambiguous_or_missing_evidence_returns_none(tmp_path: Path) -> None:
    rows = [
        _row("json", {"model_params": {"max_depth": 3}}, "a/metrics.json"),
        _row("json", {"model_params": {"max_depth": 4}}, "b/metrics.json"),
    ]
    store = _store(tmp_path, rows)
    assert AnalysisAnswerer().answer("max_depthはいくらですか", "A社", store) is None
    assert AnalysisAnswerer().answer("無関係な質問", "A社", store) is None


def test_classifier_adds_only_narrow_analysis_tags() -> None:
    assert "analysis_code" in classify_question("modeling.pyのsparse_outputがFalseになる条件は何ですか")
    assert "analysis_json" in classify_question("metrics.jsonのmodel_params.max_depthはいくらですか")
    assert "analysis_json" in classify_question("最良モデルのパラメータであるmax_depthはいくらですか")
    assert classify_question("分析結果について説明してください") == ["text_only"]
    assert "analysis_code" in classify_question("実装設定のOne-Hot Encodingカテゴリ数閾値と対象カテゴリ列を答えてください")


def test_one_hot_columns_use_effective_limit_and_complete_csv_scan(tmp_path: Path) -> None:
    import pandas as pd

    project = tmp_path / "A社" / "04.分析" / "analysis_project"
    (project / "data").mkdir(parents=True)
    pd.DataFrame({
        "small": ["a", "b", "a"], "boolean": [True, False, True],
        "numeric": [1, 2, 3], "target": [0, 1, 0],
    }).to_csv(project / "data/train.csv", index=False)
    rows = [_row("json", {
        "data_csv": "data/train.csv", "target_column": "target",
        "categorical_unique_limit_override": 3,
        "feature_plan": {"categorical_encoding": "one_hot", "categorical_unique_limit": 50},
    }, "A社/04.分析/analysis_project/configs/project_config.json")]
    answer = AnalysisAnswerer(data_dir=tmp_path).answer(
        "実装設定のOne-Hot Encodingカテゴリ数閾値を確認し、その条件で対象となるカテゴリ列をすべて答えてください", "A社", _store(tmp_path, rows)
    )
    assert answer is not None
    assert answer.text == "閾値は3未満です。対象列はsmallです。"


def test_notebook_text_output_top_n_smallest_is_verified_against_csv(tmp_path: Path) -> None:
    import pandas as pd

    project_dir = tmp_path / "A社" / "04.分析" / "analysis_project"
    (project_dir / "data").mkdir(parents=True)
    pd.DataFrame({"target": [1, 2, 3, 4, 5], "a": [1, 2, 3, 4, 5], "b": [2, 3, 4, 6, 8], "c": [5, 4, 3, 2, 1]}).to_csv(project_dir / "data/train.csv", index=False)
    rows = [
        _row("json", {"data_csv": "data/train.csv", "target_column": "target"}, "A社/04.分析/analysis_project/configs/project_config.json"),
        _row("notebook", {"cells": [{"cell_number": 1, "cell_type": "code", "source": "corr_s.abs().sort_values(ascending=False).head(3)", "outputs": [{"output_type": "stream", "text": "目的変数との相関 上位3\na  1.000000\nc  1.000000\nb  0.984798\nName: target, dtype: float64\n"}]}]}, "A社/04.分析/analysis_project/notebooks/01_eda.ipynb"),
    ]
    answer = AnalysisAnswerer(data_dir=tmp_path).answer(
        "NB01_eda.ipynbの相関 上位3の中で相関係数が最も小さいカラム名", "A社", _store(tmp_path, rows)
    )
    assert answer is not None and answer.text == "b"


def test_notebook_source_recomputes_heatmap_top_n_smallest(tmp_path: Path) -> None:
    import pandas as pd

    project_dir = tmp_path / "A社" / "04.分析" / "analysis_project"
    (project_dir / "data").mkdir(parents=True)
    pd.DataFrame({"target": [1, 2, 3, 4, 5], "a": [1, 2, 3, 4, 5], "b": [2, 3, 4, 6, 8], "c": [5, 4, 3, 2, 1]}).to_csv(project_dir / "data/train.csv", index=False)
    rows = [
        _row("json", {"data_csv": "data/train.csv", "target_column": "target"}, "A社/04.分析/analysis_project/configs/project_config.json"),
        _row("notebook", {"cells": [{"cell_number": 1, "cell_type": "code", "source": "target_corr_abs = frame.corrwith(frame[target]).abs().sort_values(ascending=False)\ntop_cols = target_corr_abs.head(3).index.tolist()\nsns.heatmap(frame[top_cols].corr())", "outputs": []}]}, "A社/04.分析/analysis_project/notebooks/01_eda.ipynb"),
    ]
    answer = AnalysisAnswerer(data_dir=tmp_path).answer(
        "01_eda.ipynbの特徴量相関ヒートマップで可視化された特徴量のうち、targetとの相関係数の絶対値が最も小さい特徴量", "A社", _store(tmp_path, rows)
    )
    assert answer is not None and answer.text == "b"


def test_notebook_correlation_returns_none_for_tie_or_axis_ticks(tmp_path: Path) -> None:
    import pandas as pd

    project_dir = tmp_path / "A社" / "04.分析" / "analysis_project"
    (project_dir / "data").mkdir(parents=True)
    pd.DataFrame({"target": [1, 2, 3], "a": [1, 2, 3], "b": [3, 2, 1]}).to_csv(project_dir / "data/train.csv", index=False)
    rows = [
        _row("json", {"data_csv": "data/train.csv", "target_column": "target"}, "A社/04.分析/analysis_project/configs/project_config.json"),
        _row("notebook", {"cells": [{"cell_number": 1, "cell_type": "code", "source": "corr.abs().sort_values(ascending=False).head(2)", "outputs": []}]}, "A社/04.分析/analysis_project/notebooks/01_eda.ipynb"),
    ]
    answerer = AnalysisAnswerer(data_dir=tmp_path)
    store = _store(tmp_path, rows)
    assert answerer.answer("01_eda.ipynbで目的変数との相関が最も高い特徴量", "A社", store) is None
    assert answerer.answer("01_eda.ipynbのy軸目盛りの最大値", "A社", store) is None


def test_heatmap_spec_is_not_contaminated_by_other_correlation_cell(tmp_path: Path) -> None:
    import pandas as pd

    project_dir = tmp_path / "A社" / "04.分析" / "analysis_project"
    (project_dir / "data").mkdir(parents=True)
    pd.DataFrame({"target": [1, 2, 3, 4, 5], "a": [1, 2, 3, 4, 5], "b": [1, 1, 2, 2, 3], "c": [5, 3, 4, 1, 2]}).to_csv(project_dir / "data/train.csv", index=False)
    rows = [
        _row("json", {"data_csv": "data/train.csv", "target_column": "target"}, "A社/04.分析/analysis_project/configs/project_config.json"),
        _row("notebook", {"cells": [
            {"cell_number": 1, "cell_type": "code", "source": "corr = frame.corr()\nsns.heatmap(corr)", "outputs": []},
            {"cell_number": 2, "cell_type": "code", "source": "corr_s.abs().sort_values(ascending=False).head(1)", "outputs": [{"output_type": "stream", "text": "目的変数との相関 上位1\na 1.0"}]},
        ]}, "A社/04.分析/analysis_project/notebooks/01_eda.ipynb"),
    ]
    answer = AnalysisAnswerer(data_dir=tmp_path).answer(
        "01_eda.ipynbの相関ヒートマップにある特徴量のうち、targetとの相関が最も小さい特徴量", "A社", _store(tmp_path, rows)
    )
    assert answer is not None and answer.text == "c"


def test_detailed_metric_improvement_uses_decimal_and_two_sources(tmp_path: Path) -> None:
    project = tmp_path / "A社"
    report = project / "05.会議" / "報告資料"
    report.mkdir(parents=True)
    from docx import Document
    doc = Document()
    doc.add_paragraph("中間報告のMacro F1: 0.7319904178115971")
    doc.save(report / "中間報告.docx")
    store = _store(tmp_path, [
        _row("json", {"f1_macro": 0.7422917255604067}, "A社/04.分析/analysis_outputs/metrics.json")
    ])
    answer = AnalysisAnswerer(data_dir=tmp_path).answer(
        "中間報告資料のMacro F1詳細値とmetrics.jsonのMacro F1詳細値を用いて改善幅を小数第6位まで", "A社", store
    )
    assert answer is not None and answer.text == "0.010301"
    assert len(answer.source_docs) == 2


def test_detailed_metric_rejects_rounded_or_multiple_report_sources(tmp_path: Path) -> None:
    from docx import Document
    report = tmp_path / "A社" / "05.会議"
    report.mkdir(parents=True)
    doc = Document(); doc.add_paragraph("Macro F1: 0.732"); doc.save(report / "a.docx")
    store = _store(tmp_path, [_row("json", {"f1_macro": 0.7422917255604067}, "metrics.json")])
    answerer = AnalysisAnswerer(data_dir=tmp_path)
    assert answerer.answer(
        "中間報告のMacro F1詳細値とmetrics.jsonの詳細値による改善幅", "A社", store
    ) is None
    doc = Document(); doc.add_paragraph("Macro F1: 0.7319904178115971"); doc.save(report / "b.docx")
    doc = Document(); doc.add_paragraph("Macro F1: 0.7300000000000000"); doc.save(report / "c.docx")
    assert answerer.answer(
        "中間報告のMacro F1詳細値とmetrics.jsonの詳細値による改善幅", "A社", store
    ) is None


def test_ranked_model_table_returns_accuracy_of_next_f1_row(tmp_path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches
    report = tmp_path / "A社" / "06.報告書"
    report.mkdir(parents=True)
    prs = Presentation(); slide = prs.slides.add_slide(prs.slide_layouts[6])
    table = slide.shapes.add_table(4, 4, Inches(1), Inches(1), Inches(8), Inches(3)).table
    rows = [
        ["Rank", "モデル種別", "F1 (macro)", "Accuracy"],
        ["1", "gradient_boosting", "0.72243", "0.89993"],
        ["2", "random_forest", "0.71486", "0.90527"],
        ["3", "linear", "0.70000", "0.88000"],
    ]
    for r, values in enumerate(rows):
        for c, value in enumerate(values): table.cell(r, c).text = value
    prs.save(report / "最終報告.pptx")
    answer = AnalysisAnswerer(data_dir=tmp_path).answer(
        "最終報告.pptxにおいて、F1スコアにてgradient_boostingに次ぐ順位のモデルのAccuracyはいくつですか", "A社", _store(tmp_path, [])
    )
    assert answer is not None and answer.text == "0.90527"


def test_ranked_model_table_rejects_missing_boundaries_or_ambiguous_files(tmp_path: Path) -> None:
    report = tmp_path / "A社" / "06.報告書"; report.mkdir(parents=True)
    (report / "最終報告.pptx").write_text("not a pptx", encoding="utf-8")
    answerer = AnalysisAnswerer(data_dir=tmp_path)
    assert answerer.answer(
        "最終報告.pptxのF1スコアでgradient_boostingに次ぐモデルのAccuracy", "A社", _store(tmp_path, [])
    ) is None
