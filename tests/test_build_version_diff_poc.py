"""pptx_lines()がテーブルシェイプ内のテキストも差分対象に含めることを確認する。

実データ（青嶺不動産アセットマネジメント案件のvalid Q9）で、旧版→新版の
「QAレビューア」変更がテーブルセル内にあり、テキストフレームのみを見る実装では
0差分に見えてしまいIncorrect回答（「変更なし」）を招いた実バグの回帰テスト。
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

from scripts.build_version_diff_poc import pptx_lines


def _make_pptx(path: Path, reviewer_name: str) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    table_shape = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(2))
    table = table_shape.table
    table.cell(0, 0).text = "QAレビューア"
    table.cell(0, 1).text = reviewer_name
    table.cell(1, 0).text = "承認日"
    table.cell(1, 1).text = "2025-08-01"
    prs.save(path)


def test_pptx_lines_includes_table_cell_text(tmp_path: Path) -> None:
    old_path = tmp_path / "old.pptx"
    new_path = tmp_path / "new.pptx"
    _make_pptx(old_path, "池田 直哉")
    _make_pptx(new_path, "小林 直樹")

    old_lines = pptx_lines(old_path)
    new_lines = pptx_lines(new_path)

    assert any("池田 直哉" in line for line in old_lines), (
        "テーブルセルのテキストがold側の抽出行に含まれていない"
    )
    assert any("小林 直樹" in line for line in new_lines), (
        "テーブルセルのテキストがnew側の抽出行に含まれていない"
    )
    assert old_lines != new_lines
