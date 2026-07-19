"""Parserコンポーネントのテスト (TDD)"""
from __future__ import annotations

import io
import json
import sys
from pathlib import Path
from unittest.mock import patch

import msoffcrypto
import pytest

import src.parsers.office_parser as office_parser_mod
from src.models import Document
from src.parsers.dispatcher import ParserDispatcher
from src.parsers.image_parser import ImageParser
from src.parsers.office_parser import OfficeParser
from src.parsers.pdf_parser import PDFParser
from src.parsers.text_parser import TextParser
from src.utils.office_crypto import derive_office_password

# ─────────────────────── TextParser ───────────────────────

class TestTextParser:
    def test_parse_txt_returns_document(self, tmp_path: Path) -> None:
        """.txtファイルをパースしてDocumentを返す"""
        f = tmp_path / "sample.txt"
        f.write_text("こんにちは、世界！", encoding="utf-8")
        parser = TextParser()
        docs = parser.parse(f)
        assert len(docs) == 1
        assert isinstance(docs[0], Document)
        assert "こんにちは" in docs[0].text

    def test_can_handle_txt(self, tmp_path: Path) -> None:
        parser = TextParser()
        assert parser.can_handle(tmp_path / "file.txt") is True

    def test_can_handle_md(self, tmp_path: Path) -> None:
        parser = TextParser()
        assert parser.can_handle(tmp_path / "file.md") is True

    def test_can_handle_csv(self, tmp_path: Path) -> None:
        parser = TextParser()
        assert parser.can_handle(tmp_path / "file.csv") is True

    def test_cannot_handle_pdf(self, tmp_path: Path) -> None:
        parser = TextParser()
        assert parser.can_handle(tmp_path / "file.pdf") is False

    def test_can_handle_json(self, tmp_path: Path) -> None:
        parser = TextParser()
        assert parser.can_handle(tmp_path / "metrics.json") is True

    def test_can_handle_tsv(self, tmp_path: Path) -> None:
        parser = TextParser()
        assert parser.can_handle(tmp_path / "data.tsv") is True

    def test_can_handle_py(self, tmp_path: Path) -> None:
        """.pyファイル（modeling.py等）が未対応だと、code_static設問の検索対象が
        存在しないまま何一つ拾えない（実測: valid Q4/Q28がともに検索失敗）。"""
        parser = TextParser()
        assert parser.can_handle(tmp_path / "modeling.py") is True

    def test_parse_py_returns_document_with_source_text(self, tmp_path: Path) -> None:
        f = tmp_path / "modeling.py"
        f.write_text("if df['CAT'].dtype == 'object':\n    pass\n", encoding="utf-8")
        parser = TextParser()
        docs = parser.parse(f)
        assert len(docs) == 1
        assert "df['CAT'].dtype" in docs[0].text


# ─────────────────────── PDFParser ───────────────────────

class TestPDFParser:
    def test_can_handle_pdf(self, tmp_path: Path) -> None:
        parser = PDFParser()
        assert parser.can_handle(tmp_path / "file.pdf") is True

    def test_cannot_handle_txt(self, tmp_path: Path) -> None:
        parser = PDFParser()
        assert parser.can_handle(tmp_path / "file.txt") is False

    def test_stub_when_pypdf_not_installed(self, tmp_path: Path) -> None:
        """pypdf未インストール時にスタブDocumentを返す（文字列に '[PDF未解析' が含まれる）"""
        f = tmp_path / "sample.pdf"
        f.write_bytes(b"%PDF-1.4 stub")
        parser = PDFParser()
        # sys.modules に None を設定すると ImportError 扱いになる
        with patch.dict(sys.modules, {"pypdf": None}):
            docs = parser.parse(f)
        assert len(docs) == 1
        assert "[PDF未解析" in docs[0].text

    def test_pdf_parse_failure_returns_stub_not_raises(self, tmp_path: Path) -> None:
        """暗号化・破損したPDFでも例外を投げずスタブを返す"""
        f = tmp_path / "broken.pdf"
        f.write_bytes(b"%PDF-1.4 this is not a valid pdf structure")
        parser = PDFParser()
        docs = parser.parse(f)
        assert len(docs) == 1
        assert "解析失敗" in docs[0].text or "empty" == docs[0].location


# ─────────────────────── OfficeParser ───────────────────────

# msoffcrypto-tool 6.0.0のOLEコンテナ書き込みは、暗号化payloadが小さい(<=4096バイト)と
# mini-FAT/regular-FATの不整合で壊れることがある（tests/test_office_crypto.pyで裏取り済み）。
# xlsxは十分な行数を足して4KB超のペイロードにする。
def _build_padded_xlsx_bytes() -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "スケジュール"
    ws.append(["タスクID", "担当者", "開始日"])
    for i in range(300):
        ws.append([f"T{i:03d}", "テスト担当者", f"padding padding padding padding {i}"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _build_padded_docx_bytes(marker_text: str) -> bytes:
    import docx

    doc = docx.Document()
    doc.add_paragraph(marker_text)
    for i in range(50):
        doc.add_paragraph(f"padding paragraph {i} " * 10)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _write_encrypted_office_file(dest: Path, password: str, plaintext_bytes: bytes) -> None:
    plain_buf = io.BytesIO(plaintext_bytes)
    office_file = msoffcrypto.OfficeFile(plain_buf)
    dest.parent.mkdir(parents=True, exist_ok=True)
    with open(dest, "wb") as f:
        office_file.encrypt(password, f)


def _write_registry_and_contracts(
    tmp_path: Path, project_name: str, alias: str, start_date: str
) -> tuple[Path, Path]:
    registry_path = tmp_path / "project_registry.json"
    registry_path.write_text(
        json.dumps(
            [{"project_name": project_name, "primary_alias": alias}], ensure_ascii=False
        ),
        encoding="utf-8",
    )
    contracts_path = tmp_path / "contracts.jsonl"
    contracts_path.write_text(
        json.dumps(
            {"project_name": project_name, "start_date": start_date}, ensure_ascii=False
        )
        + "\n",
        encoding="utf-8",
    )
    return registry_path, contracts_path


class TestOfficeParserEncryptedContentDetection:
    """暗号化検知をファイル名`pw-`規則からファイル内容(CDFV2/OLEマジックナンバー)ベースへ
    一般化する回帰テスト。KAEDEのスケジュール.xlsx（ファイル名に`pw-`マーカーが一切無い
    CDFV2暗号化ファイル）が復号されずインデックスから完全欠落していた問題(Q79)に対応する。
    """

    def test_xlsx_decrypts_via_content_signature_without_pw_filename_marker(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """ファイル名に`pw-`が無くてもCDFV2署名だけで暗号化と判定し、DA-規則パスワード
        (案件レジストリのprimary_alias + contracts.jsonlの開始日)を導出して復号できる。"""
        project_dir = tmp_path / "プロジェクト" / "テスト案件" / "02.計画"
        xlsx_path = project_dir / "スケジュール.xlsx"
        password = derive_office_password("KAEDE", "2025-09-02", ".xlsx")
        _write_encrypted_office_file(xlsx_path, password, _build_padded_xlsx_bytes())

        registry_path, contracts_path = _write_registry_and_contracts(
            tmp_path, "テスト案件", "KAEDE", "2025-09-02"
        )
        monkeypatch.setattr(office_parser_mod, "PROJECT_REGISTRY_PATH", registry_path)
        monkeypatch.setattr(office_parser_mod, "CONTRACTS_PATH", contracts_path)

        docs = OfficeParser().parse(xlsx_path)

        assert not any("解析失敗" in d.text for d in docs)
        assert any("テスト担当者" in d.text for d in docs)

    def test_docx_decrypts_via_content_signature_without_pw_filename_marker(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """xlsxと同じ内容ベース検知・DA-規則復号がdocxにも汎用的に効くことの確認。"""
        project_dir = tmp_path / "プロジェクト" / "テスト案件" / "05.会議"
        docx_path = project_dir / "会議録.docx"
        password = derive_office_password("KAEDE", "2025-09-02", ".docx")
        _write_encrypted_office_file(
            docx_path, password, _build_padded_docx_bytes("会議の要点マーカーテキスト")
        )

        registry_path, contracts_path = _write_registry_and_contracts(
            tmp_path, "テスト案件", "KAEDE", "2025-09-02"
        )
        monkeypatch.setattr(office_parser_mod, "PROJECT_REGISTRY_PATH", registry_path)
        monkeypatch.setattr(office_parser_mod, "CONTRACTS_PATH", contracts_path)

        docs = OfficeParser().parse(docx_path)

        assert not any("解析失敗" in d.text for d in docs)
        assert any("会議の要点マーカーテキスト" in d.text for d in docs)

    def test_xlsx_still_decrypts_via_pw_filename_literal_password(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """既存の`pw-<トークン>`命名慣習（リテラルパスワード）が退行していないことの確認。"""
        xlsx_path = tmp_path / "スケジュール_pw-testtoken123.xlsx"
        _write_encrypted_office_file(xlsx_path, "testtoken123", _build_padded_xlsx_bytes())

        # 案件レジストリ/contracts側は存在しなくてもリテラル候補だけで復号できること
        monkeypatch.setattr(
            office_parser_mod, "PROJECT_REGISTRY_PATH", tmp_path / "no_registry.json"
        )
        monkeypatch.setattr(office_parser_mod, "CONTRACTS_PATH", tmp_path / "no_contracts.jsonl")

        docs = OfficeParser().parse(xlsx_path)

        assert not any("解析失敗" in d.text for d in docs)
        assert any("テスト担当者" in d.text for d in docs)

    def test_docx_pw_filename_falls_back_from_literal_to_da_rule_password(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """実データの契約書_pw-kaede20250902.docx型の命名: ファイル名の`pw-`トークンを
        素直にリテラルパスワードとして試すと失敗するが（実際の暗号化パスワードはDA-規則側）、
        同じマーカーから抽出した日付でDA-規則パスワードにフォールバックして復号できること。
        """
        project_dir = tmp_path / "プロジェクト" / "テスト案件" / "01.契約"
        docx_path = project_dir / "契約書_pw-kaede20250902.docx"
        da_rule_password = derive_office_password("KAEDE", "20250902", ".docx")
        _write_encrypted_office_file(
            docx_path, da_rule_password, _build_padded_docx_bytes("契約条件マーカーテキスト")
        )

        registry_path, contracts_path = _write_registry_and_contracts(
            tmp_path, "テスト案件", "KAEDE", "2025-09-02"
        )
        monkeypatch.setattr(office_parser_mod, "PROJECT_REGISTRY_PATH", registry_path)
        monkeypatch.setattr(office_parser_mod, "CONTRACTS_PATH", contracts_path)

        docs = OfficeParser().parse(docx_path)

        assert not any("解析失敗" in d.text for d in docs)
        assert any("契約条件マーカーテキスト" in d.text for d in docs)

    def test_xlsx_decrypt_falls_back_to_stub_when_no_password_candidate_matches(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """CDFV2署名はあるが導出できるどの候補パスワードでも復号できない場合、
        例外を出さず既存のstub Documentへフォールバックすること。"""
        project_dir = tmp_path / "プロジェクト" / "テスト案件" / "02.計画"
        xlsx_path = project_dir / "スケジュール.xlsx"
        real_password = derive_office_password("REAL", "2025-01-01", ".xlsx")
        _write_encrypted_office_file(xlsx_path, real_password, _build_padded_xlsx_bytes())

        registry_path, contracts_path = _write_registry_and_contracts(
            tmp_path, "テスト案件", "WRONG", "2025-09-02"
        )
        monkeypatch.setattr(office_parser_mod, "PROJECT_REGISTRY_PATH", registry_path)
        monkeypatch.setattr(office_parser_mod, "CONTRACTS_PATH", contracts_path)

        docs = OfficeParser().parse(xlsx_path)

        assert len(docs) == 1
        assert "解析失敗" in docs[0].text

    def test_xlsx_generic_decryption_error_on_one_candidate_still_tries_next(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """msoffcryptoが`InvalidKeyError`(パスワード誤り)ではなくその親クラス
        `DecryptionError`を送出するケース（暗号化情報ストリームの解釈失敗等）でも、
        その候補で例外を伝播させず次の候補に進めること。"""
        from msoffcrypto.exceptions import DecryptionError

        # 候補を2つ以上用意する: (1)ファイル名の`pw-`リテラル候補（実際の正解ではない）を
        # 1回目のDecryptionErrorで潰し、(2)DA規則候補（実際の正解）を2回目で使う。
        project_dir = tmp_path / "プロジェクト" / "テスト案件" / "02.計画"
        xlsx_path = project_dir / "スケジュール_pw-irrelevanttoken.xlsx"
        password = derive_office_password("KAEDE", "2025-09-02", ".xlsx")
        _write_encrypted_office_file(xlsx_path, password, _build_padded_xlsx_bytes())

        registry_path, contracts_path = _write_registry_and_contracts(
            tmp_path, "テスト案件", "KAEDE", "2025-09-02"
        )
        monkeypatch.setattr(office_parser_mod, "PROJECT_REGISTRY_PATH", registry_path)
        monkeypatch.setattr(office_parser_mod, "CONTRACTS_PATH", contracts_path)

        real_decrypt = office_parser_mod.decrypt_office_file
        calls: list[str] = []

        def flaky_decrypt(path, pw, output_path):
            calls.append(pw)
            if len(calls) == 1:
                raise DecryptionError("simulated non-InvalidKeyError decryption failure")
            return real_decrypt(path, pw, output_path)

        monkeypatch.setattr(office_parser_mod, "decrypt_office_file", flaky_decrypt)

        docs = OfficeParser().parse(xlsx_path)

        assert len(calls) >= 2, "1回目のDecryptionErrorで打ち切らず次の候補を試していない"
        assert not any("解析失敗" in d.text for d in docs)
        assert any("テスト担当者" in d.text for d in docs)

    def test_xlsx_non_cdfv2_corruption_does_not_attempt_decrypt(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """CDFV2署名の無い単なる破損ファイルは、復号を試みず従来どおりstubを返すこと
        （暗号化以外の破損原因まで復号フォールバックの対象にしないための回帰確認）。"""
        xlsx_path = tmp_path / "プロジェクト" / "テスト案件" / "02.計画" / "スケジュール.xlsx"
        xlsx_path.parent.mkdir(parents=True)
        xlsx_path.write_bytes(b"not a real xlsx file, just garbage bytes")

        registry_path, contracts_path = _write_registry_and_contracts(
            tmp_path, "テスト案件", "KAEDE", "2025-09-02"
        )
        monkeypatch.setattr(office_parser_mod, "PROJECT_REGISTRY_PATH", registry_path)
        monkeypatch.setattr(office_parser_mod, "CONTRACTS_PATH", contracts_path)

        docs = OfficeParser().parse(xlsx_path)

        assert len(docs) == 1
        assert "解析失敗" in docs[0].text


class TestOfficeParser:
    def test_can_handle_docx(self, tmp_path: Path) -> None:
        parser = OfficeParser()
        assert parser.can_handle(tmp_path / "file.docx") is True

    def test_can_handle_xlsx(self, tmp_path: Path) -> None:
        parser = OfficeParser()
        assert parser.can_handle(tmp_path / "file.xlsx") is True

    def test_can_handle_pptx(self, tmp_path: Path) -> None:
        parser = OfficeParser()
        assert parser.can_handle(tmp_path / "file.pptx") is True

    def test_cannot_handle_pdf(self, tmp_path: Path) -> None:
        parser = OfficeParser()
        assert parser.can_handle(tmp_path / "file.pdf") is False

    def test_stub_docx_when_not_installed(self, tmp_path: Path) -> None:
        """python-docx未インストール時にスタブ文字列を含むDocumentを返す"""
        f = tmp_path / "sample.docx"
        f.write_bytes(b"PK stub")
        parser = OfficeParser()
        with patch.dict(sys.modules, {"docx": None}):
            docs = parser.parse(f)
        assert len(docs) == 1
        assert "DOCX未解析" in docs[0].text

    def test_stub_xlsx_when_not_installed(self, tmp_path: Path) -> None:
        """openpyxl未インストール時にスタブ文字列を含むDocumentを返す"""
        f = tmp_path / "sample.xlsx"
        f.write_bytes(b"PK stub")
        parser = OfficeParser()
        with patch.dict(sys.modules, {"openpyxl": None}):
            docs = parser.parse(f)
        assert len(docs) == 1
        assert "XLSX未解析" in docs[0].text

    def test_stub_pptx_when_not_installed(self, tmp_path: Path) -> None:
        """python-pptx未インストール時にスタブ文字列を含むDocumentを返す"""
        f = tmp_path / "sample.pptx"
        f.write_bytes(b"PK stub")
        parser = OfficeParser()
        with patch.dict(sys.modules, {"pptx": None}):
            docs = parser.parse(f)
        assert len(docs) == 1
        assert "PPTX未解析" in docs[0].text

    def test_docx_parse_failure_returns_stub_not_raises(self, tmp_path: Path) -> None:
        """壊れた/パスワード保護されたdocxでも例外を投げずスタブを返す"""
        f = tmp_path / "broken.docx"
        f.write_bytes(b"not a real docx file, definitely not a valid zip")
        parser = OfficeParser()
        docs = parser.parse(f)
        assert len(docs) == 1
        assert "解析失敗" in docs[0].text

    def test_docx_parse_includes_table_cell_text(self, tmp_path: Path) -> None:
        """Wordの表（doc.tables）のセルテキストも索引対象に含める。

        現状は doc.paragraphs のみを結合しており doc.tables を無視しているため、
        表形式の内容（略語辞典・契約条件表等）がまるごと検索対象から消える
        （実データ: 社内用語集.docxが9表820セルのうち0セルしか索引されない）。
        """
        import docx

        f = tmp_path / "with_table.docx"
        doc = docx.Document()
        doc.add_paragraph("見出し")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "略語"
        table.cell(0, 1).text = "意味"
        table.cell(1, 0).text = "FM"
        table.cell(1, 1).text = "ファシリティマネジメント"
        doc.save(f)

        parser = OfficeParser()
        docs = parser.parse(f)

        assert any("ファシリティマネジメント" in d.text for d in docs), (
            "表セルのテキストがパース結果に含まれていない"
        )

    def test_docx_parse_preserves_body_order_of_paragraphs_and_tables(
        self, tmp_path: Path
    ) -> None:
        """段落と表が混在する文書で、本文中の出現順を保持する。"""
        import docx

        f = tmp_path / "ordered.docx"
        doc = docx.Document()
        doc.add_paragraph("最初の段落")
        table = doc.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "表の内容"
        doc.add_paragraph("最後の段落")
        doc.save(f)

        parser = OfficeParser()
        docs = parser.parse(f)
        text = "\n".join(d.text for d in docs)

        first_idx = text.index("最初の段落")
        table_idx = text.index("表の内容")
        last_idx = text.index("最後の段落")
        assert first_idx < table_idx < last_idx, (
            "段落と表の出現順（body順）が保持されていない"
        )

    def test_docx_parse_handles_merged_and_empty_cells_without_raising(
        self, tmp_path: Path
    ) -> None:
        """結合セル・空セルを含む表でも例外を出さず、非空セルの内容を抽出する。

        python-docxの`row.cells`は横結合(gridSpan)されたセルをスパン列数分
        同一tc要素として複数回yieldする仕様のため、そのまま連結すると
        「結合セル見出し | 結合セル見出し」のように内容が重複してしまう。
        tc要素の同一性でデデュープし、重複なく1回だけ抽出されることを検証する。
        """
        import docx

        f = tmp_path / "merged_cells.docx"
        doc = docx.Document()
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).merge(table.cell(0, 1))
        table.cell(0, 0).text = "結合セル見出し"
        table.cell(1, 0).text = ""
        table.cell(1, 1).text = "右下セルの値"
        doc.save(f)

        parser = OfficeParser()
        docs = parser.parse(f)  # 例外を出さないこと

        text = "\n".join(d.text for d in docs)
        assert text.count("結合セル見出し") == 1, (
            "横結合セルの内容が重複して抽出されている: " + repr(text)
        )
        assert "右下セルの値" in text

    def test_docx_parse_dedups_vertically_merged_cell_across_rows(
        self, tmp_path: Path
    ) -> None:
        """縦結合(vMerge)セルの継続行でも、python-docxは起点セルの内容を
        再度yieldする仕様のため、そのまま連結すると起点セルの内容が
        行数分だけ丸ごと重複してしまう。tc要素の同一性で表全体にわたって
        デデュープし、1回だけ抽出されることを検証する。
        """
        import docx

        f = tmp_path / "vmerged.docx"
        doc = docx.Document()
        table = doc.add_table(rows=3, cols=2)
        table.cell(0, 0).merge(table.cell(1, 0))
        table.cell(0, 0).text = "縦結合の内容"
        table.cell(0, 1).text = "上段の値"
        table.cell(1, 1).text = "下段の値"
        table.cell(2, 0).text = "無関係セル"
        table.cell(2, 1).text = "無関係の値"
        doc.save(f)

        parser = OfficeParser()
        docs = parser.parse(f)

        text = "\n".join(d.text for d in docs)
        assert text.count("縦結合の内容") == 1, (
            "縦結合セルの内容が行ごとに重複して抽出されている: " + repr(text)
        )
        assert "上段の値" in text
        assert "下段の値" in text
        assert "無関係セル" in text

    def test_docx_parse_handles_nested_table_without_raising(self, tmp_path: Path) -> None:
        """表のセル内にネストした表があっても例外を出さない。"""
        import docx

        f = tmp_path / "nested_table.docx"
        doc = docx.Document()
        outer_table = doc.add_table(rows=1, cols=1)
        outer_cell = outer_table.cell(0, 0)
        outer_cell.text = "外側セル"
        nested_table = outer_cell.add_table(rows=1, cols=1)
        nested_table.cell(0, 0).text = "内側セル"
        doc.save(f)

        parser = OfficeParser()
        docs = parser.parse(f)  # 例外を出さないこと

        assert len(docs) == 1
        assert "外側セル" in docs[0].text

    def test_xlsx_parse_failure_returns_stub_not_raises(self, tmp_path: Path) -> None:
        f = tmp_path / "broken.xlsx"
        f.write_bytes(b"not a real xlsx file")
        parser = OfficeParser()
        docs = parser.parse(f)
        assert len(docs) == 1
        assert "解析失敗" in docs[0].text

    def test_pptx_parse_failure_returns_stub_not_raises(self, tmp_path: Path) -> None:
        f = tmp_path / "broken.pptx"
        f.write_bytes(b"not a real pptx file")
        parser = OfficeParser()
        docs = parser.parse(f)
        assert len(docs) == 1
        assert "解析失敗" in docs[0].text

    def test_pptx_parse_includes_table_cell_text(self, tmp_path: Path) -> None:
        """テーブルシェイプ（GraphicFrame）内のセルテキストも索引対象に含める。

        GraphicFrameは.text属性を持たないため、通常のテキストフレームのみを見る
        実装だとスライド内の表の内容（例: レビューア・承認欄）が丸ごと検索対象から
        消える（実データ: 青嶺不動産アセットマネジメント案件のQ9でIncorrectを誘発）。
        """
        from pptx import Presentation
        from pptx.util import Inches

        f = tmp_path / "with_table.pptx"
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        table_shape = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(2))
        table = table_shape.table
        table.cell(0, 0).text = "QAレビューア"
        table.cell(0, 1).text = "小林 直樹"
        prs.save(f)

        parser = OfficeParser()
        docs = parser.parse(f)

        assert any("小林 直樹" in doc.text for doc in docs), (
            "テーブルセルのテキストがパース結果に含まれていない"
        )


# ─────────────────────── ImageParser ───────────────────────

class TestImageParser:
    def test_can_handle_png(self, tmp_path: Path) -> None:
        parser = ImageParser()
        assert parser.can_handle(tmp_path / "file.png") is True

    def test_can_handle_jpg(self, tmp_path: Path) -> None:
        parser = ImageParser()
        assert parser.can_handle(tmp_path / "file.jpg") is True

    def test_can_handle_jpeg(self, tmp_path: Path) -> None:
        parser = ImageParser()
        assert parser.can_handle(tmp_path / "file.jpeg") is True

    def test_parse_returns_stub_document(self, tmp_path: Path) -> None:
        """parse() がスタブ文字列を含む Document を返す"""
        f = tmp_path / "sample.png"
        f.write_bytes(b"\x89PNG stub")
        parser = ImageParser()
        docs = parser.parse(f)
        assert len(docs) == 1
        assert isinstance(docs[0], Document)
        # スタブ実装は「画像」または「VLM」を含む文字列を返す
        assert "画像" in docs[0].text or "VLM" in docs[0].text

    def test_metadata_contains_type_image(self, tmp_path: Path) -> None:
        """metadata に type: image が含まれる"""
        f = tmp_path / "photo.jpg"
        f.write_bytes(b"JFIF stub")
        parser = ImageParser()
        docs = parser.parse(f)
        assert docs[0].metadata.get("type") == "image"


# ─────────────────────── ParserDispatcher ───────────────────────

class TestParserDispatcher:
    def test_dispatch_txt_file(self, tmp_path: Path) -> None:
        """.txtファイルをTextParserでパースする（Document.text が空でない）"""
        f = tmp_path / "hello.txt"
        f.write_text("テキストデータ", encoding="utf-8")
        dispatcher = ParserDispatcher()
        docs = dispatcher.parse(f)
        assert len(docs) >= 1
        assert docs[0].text != ""
        assert "テキストデータ" in docs[0].text

    def test_unsupported_extension_returns_stub(self, tmp_path: Path) -> None:
        """未対応拡張子のファイルを渡すと '[未対応形式' を含むDocumentを返す"""
        f = tmp_path / "weird.xyz"
        f.write_bytes(b"some data")
        dispatcher = ParserDispatcher()
        docs = dispatcher.parse(f)
        assert len(docs) == 1
        assert "[未対応形式" in docs[0].text

    def test_parse_directory_recursive(self, tmp_path: Path) -> None:
        """parse_directory() でディレクトリ内の複数ファイルを再帰的に処理できる"""
        (tmp_path / "a.txt").write_text("ファイルA", encoding="utf-8")
        (tmp_path / "b.md").write_text("ファイルB", encoding="utf-8")
        sub = tmp_path / "sub"
        sub.mkdir()
        (sub / "c.txt").write_text("ファイルC", encoding="utf-8")

        dispatcher = ParserDispatcher()
        docs = dispatcher.parse_directory(tmp_path)
        texts = [d.text for d in docs]
        assert any("ファイルA" in t for t in texts)
        assert any("ファイルB" in t for t in texts)
        assert any("ファイルC" in t for t in texts)

    def test_parse_directory_extracts_project_metadata(self, tmp_path: Path) -> None:
        """プロジェクト/<企業名>/<カテゴリ>/ 配下のファイルにproject/categoryメタデータを付与する"""
        proj_dir = tmp_path / "プロジェクト" / "サンプル株式会社" / "00.提案"
        proj_dir.mkdir(parents=True)
        (proj_dir / "doc.txt").write_text("提案内容", encoding="utf-8")

        dispatcher = ParserDispatcher()
        docs = dispatcher.parse_directory(tmp_path)

        assert len(docs) == 1
        assert docs[0].metadata["project"] == "サンプル株式会社"
        assert docs[0].metadata["category"] == "00.提案"
        assert docs[0].metadata["is_internal"] is False

    def test_parse_directory_extracts_project_metadata_from_nfd_paths(self, tmp_path: Path) -> None:
        """macOSのファイルシステム／zip展開由来のNFDパス（「プロジェクト」の「ジ」が
        分解されている等）でもproject/categoryメタデータを付与できる。
        これが失敗すると実データで全ドキュメントのprojectがNoneになり、
        案件スコープ検索・構造化ルーティングが一切機能しなくなる。"""
        import unicodedata

        nfd_projects = unicodedata.normalize("NFD", "プロジェクト")
        nfd_company = unicodedata.normalize("NFD", "青葉与信マネジメント株式会社")
        assert nfd_projects != "プロジェクト"  # 前提: 分解可能な文字を含む

        proj_dir = tmp_path / nfd_projects / nfd_company / "02.計画"
        proj_dir.mkdir(parents=True)
        (proj_dir / "doc.txt").write_text("計画内容", encoding="utf-8")

        dispatcher = ParserDispatcher()
        docs = dispatcher.parse_directory(tmp_path)

        assert len(docs) == 1
        assert docs[0].metadata["project"] == "青葉与信マネジメント株式会社"  # NFCで返る
        assert docs[0].metadata["category"] == "02.計画"

    def test_parse_directory_marks_internal_docs(self, tmp_path: Path) -> None:
        """社内管理/ 配下のファイルは is_internal=True, project=None になる"""
        internal_dir = tmp_path / "社内管理"
        internal_dir.mkdir()
        (internal_dir / "用語集.txt").write_text("用語", encoding="utf-8")

        dispatcher = ParserDispatcher()
        docs = dispatcher.parse_directory(tmp_path)

        assert docs[0].metadata["is_internal"] is True
        assert docs[0].metadata["project"] is None

    def test_parse_directory_indexes_python_source_files(self, tmp_path: Path) -> None:
        """modeling.py等の.pyファイルが検索対象として拾われることを保証する回帰テスト
        （実測: valid Q4/Q28はcode_staticタイプで.pyが未対応のため検索失敗していた）。"""
        (tmp_path / "modeling.py").write_text(
            "if df['CAT'].dtype == 'object' and df['CAT'].nunique() < 10:\n"
            "    category_columns.append('CAT')\n",
            encoding="utf-8",
        )
        dispatcher = ParserDispatcher()
        docs = dispatcher.parse_directory(tmp_path)
        texts = [d.text for d in docs]
        assert any("df['CAT'].dtype" in t for t in texts)
        assert not any("未対応形式" in t for t in texts)

    def test_parse_directory_skips_noise_files(self, tmp_path: Path) -> None:
        """__pycache__/*.pyc, *.lock, ~$で始まる一時ファイルを除外する"""
        (tmp_path / "a.txt").write_text("残す", encoding="utf-8")
        cache_dir = tmp_path / "__pycache__"
        cache_dir.mkdir()
        (cache_dir / "mod.cpython-311.pyc").write_bytes(b"stub")
        (tmp_path / "uv.lock").write_text("lock", encoding="utf-8")
        (tmp_path / "~$temp.xlsx").write_bytes(b"lock")

        dispatcher = ParserDispatcher()
        docs = dispatcher.parse_directory(tmp_path)

        texts = [d.text for d in docs]
        assert len(docs) == 1
        assert any("残す" in t for t in texts)


# ─────────────────────── NotebookParser ───────────────────────

class TestNotebookParser:
    def test_can_handle_ipynb(self, tmp_path: Path) -> None:
        from src.parsers.notebook_parser import NotebookParser
        parser = NotebookParser()
        assert parser.can_handle(tmp_path / "01_eda.ipynb") is True

    def test_cannot_handle_txt(self, tmp_path: Path) -> None:
        from src.parsers.notebook_parser import NotebookParser
        parser = NotebookParser()
        assert parser.can_handle(tmp_path / "file.txt") is False

    def test_extracts_markdown_and_code_cell_text(self, tmp_path: Path) -> None:
        import json as jsonlib

        from src.parsers.notebook_parser import NotebookParser

        notebook = {
            "cells": [
                {"cell_type": "markdown", "source": ["# EDA\n", "欠損値を確認する"]},
                {"cell_type": "code", "source": ["df.isnull().sum()"]},
                {"cell_type": "code", "source": [""]},
            ]
        }
        f = tmp_path / "01_eda.ipynb"
        f.write_text(jsonlib.dumps(notebook), encoding="utf-8")

        parser = NotebookParser()
        docs = parser.parse(f)

        assert len(docs) == 2  # 空セルは除外される
        assert any("欠損値を確認する" in d.text for d in docs)
        assert any("df.isnull().sum()" in d.text for d in docs)

    def test_invalid_notebook_returns_stub(self, tmp_path: Path) -> None:
        from src.parsers.notebook_parser import NotebookParser

        f = tmp_path / "broken.ipynb"
        f.write_text("not valid json {{{", encoding="utf-8")

        parser = NotebookParser()
        docs = parser.parse(f)

        assert len(docs) == 1
        assert "解析失敗" in docs[0].text


# ─────────────────────── dispatcherへのNotebook登録 ───────────────────────

class TestDispatcherHandlesNotebook:
    def test_dispatch_ipynb_file(self, tmp_path: Path) -> None:
        import json as jsonlib
        notebook = {"cells": [{"cell_type": "markdown", "source": ["EDAメモ"]}]}
        f = tmp_path / "01_eda.ipynb"
        f.write_text(jsonlib.dumps(notebook), encoding="utf-8")

        dispatcher = ParserDispatcher()
        docs = dispatcher.parse(f)
        assert any("EDAメモ" in d.text for d in docs)


def test_parse_directory_excludes_dirs(tmp_path):
    from src.parsers.dispatcher import ParserDispatcher

    data_dir = tmp_path / "data"
    data_dir.mkdir()
    (data_dir / "a.txt").write_text("コーパス文書", encoding="utf-8")
    qa_dir = data_dir / "質問回答"
    qa_dir.mkdir()
    (qa_dir / "questions_valid.csv").write_text(
        "index,question,answer\n0,テスト質問,テスト正解\n", encoding="utf-8"
    )
    docs = ParserDispatcher().parse_directory(data_dir, exclude_dirs=[qa_dir])
    assert any("コーパス文書" in d.text for d in docs)
    assert not any("questions_valid" in str(d.source_path) for d in docs)


def test_parse_directory_excludes_nfd_dir_with_nfc_argument(tmp_path):
    # macOS/zip展開由来のNFDディレクトリ名を、NFCで指定したexclude_dirsで除外できること
    # （_extract_metadataと同じNFC事故パターンの回帰テスト）
    import unicodedata
    from src.parsers.dispatcher import ParserDispatcher

    data_dir = tmp_path / "data"
    data_dir.mkdir()
    (data_dir / "a.txt").write_text("コーパス文書", encoding="utf-8")
    nfd_name = unicodedata.normalize("NFD", "データ質問")
    qa_dir = data_dir / nfd_name
    qa_dir.mkdir()
    (qa_dir / "questions_valid.csv").write_text(
        "index,question,answer\n0,テスト質問,テスト正解\n", encoding="utf-8"
    )
    nfc_dir = data_dir / unicodedata.normalize("NFC", "データ質問")
    docs = ParserDispatcher().parse_directory(data_dir, exclude_dirs=[nfc_dir])
    assert any("コーパス文書" in d.text for d in docs)
    assert not any("questions_valid" in str(d.source_path) for d in docs)
